# -*- coding: utf-8 -*-
"""V6.16：送外审图表期刊化重绘；不修改实验数据。"""
from pathlib import Path
import json

import matplotlib.pyplot as plt
import numpy as np
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = Path(r"D:\IdeaProject\faladi\three_phase_ga_optimizer")
SOURCE = ROOT / "论文" / "目录约束双档案遗传优化三相变压器_版式与引用优化版V6.11-送外审正文对齐终校版.docx"
OUTPUT = ROOT / "论文" / "目录约束双档案遗传优化三相变压器_版式与引用优化版V6.16-送外审图表可读性修订版.docx"
ASSET = ROOT / "论文" / "图表-送外审终校V6.16"
DATA = ROOT / "测试归档" / "2026-09-24测试" / "论文真实数据"

INK, BLUE, TEAL, ORANGE, GRID = "#19324D", "#2E5F89", "#167C74", "#BE7A2F", "#D9E3EC"
PALE_BLUE, PALE_TEAL, PALE_ORANGE = "#F4F8FC", "#EEF8F6", "#FFF7EE"
plt.rcParams.update({"font.sans-serif": ["Microsoft YaHei", "SimHei", "Arial Unicode MS"], "axes.unicode_minus": False})


def save(fig, path):
    fig.savefig(path, dpi=600, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def arrow(ax, a, b, color=BLUE, lw=1.8, rad=0):
    ax.add_patch(FancyArrowPatch(a, b, arrowstyle="-|>", mutation_scale=12,
                                 linewidth=lw, color=color, connectionstyle=f"arc3,rad={rad}"))


def box(ax, x, y, w, h, title, note="", fill=PALE_BLUE, edge=BLUE, ts=13, ns=10):
    ax.add_patch(FancyBboxPatch((x,y), w,h, boxstyle="round,pad=0.02,rounding_size=0.035",
                                facecolor=fill, edgecolor=edge, linewidth=1.55))
    ax.text(x+w/2,y+h*.61,title,ha="center",va="center",fontsize=ts,fontweight="bold",color=INK)
    if note: ax.text(x+w/2,y+h*.28,note,ha="center",va="center",fontsize=ns,color="#4D6278")


def paper_box(ax, x, y, w, h, text, fs=13):
    """单色直角框：用于期刊流程示意，避免演示文稿式配色。"""
    ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle="square,pad=0.01",
                                facecolor="white", edgecolor="#3D4C5C", linewidth=1.05))
    ax.text(x+w/2, y+h/2, text, ha="center", va="center", fontsize=fs,
            fontweight="normal", color="#1F2933")


def fig1(path):
    # 提高图高并缩短框内文字，避免双栏缩放后文字或箭头贴边。
    fig,ax=plt.subplots(figsize=(6.2,3.25)); ax.set_xlim(0,12); ax.set_ylim(0,6); ax.axis("off")
    c="#3D4C5C"
    paper_box(ax,.28,4.45,2.55,.92,"页面约束\n冻结目录",12.6)
    paper_box(ax,3.55,4.45,1.90,.92,"合法化",13.5)
    paper_box(ax,6.25,4.45,1.98,.92,"工程精算",13.5)
    paper_box(ax,9.02,4.45,2.35,.92,"约束诊断",13.5)
    paper_box(ax,6.25,1.50,1.98,.92,"双档案\nA_f / A_n",12.7)
    paper_box(ax,3.55,1.50,1.90,.92,"搜索更新",13.5)
    for a,b in [((2.83,4.91),(3.55,4.91)),((5.45,4.91),(6.25,4.91)),((8.23,4.91),(9.02,4.91)),((10.20,4.45),(10.20,2.42)),((9.02,1.96),(8.23,1.96)),((6.25,1.96),(5.45,1.96)),((4.50,2.42),(4.50,4.45))]:
        arrow(ax,a,b,c,lw=1.15)
    fig.tight_layout(pad=.12); save(fig,path)


def fig2(path):
    # 只保留论证所需的两类操作，增大高度和字号以保证单栏可读。
    fig,ax=plt.subplots(figsize=(6.2,4.25)); ax.set_xlim(0,12); ax.set_ylim(0,8); ax.axis("off")
    c="#3D4C5C"
    paper_box(ax,.55,5.55,2.55,1.00,"目录记录 A\n完整字段",16.0)
    paper_box(ax,4.10,5.55,3.80,1.00,"字段拼接：禁止\n目录不可还原",16.0)
    paper_box(ax,8.90,5.55,2.55,1.00,"目录记录 B\n完整字段",16.0)
    arrow(ax,(3.10,6.05),(4.10,6.05),c,lw=1.12); arrow(ax,(8.90,6.05),(7.90,6.05),c,lw=1.12)
    paper_box(ax,.72,1.65,4.25,1.05,"完整记录整体继承",16.4)
    paper_box(ax,5.30,1.65,2.10,1.05,"合法化",16.4)
    paper_box(ax,7.73,1.65,3.55,1.05,"工程精算",16.4)
    arrow(ax,(4.97,2.18),(5.30,2.18),c,lw=1.12); arrow(ax,(7.40,2.18),(7.73,2.18),c,lw=1.12)
    fig.tight_layout(pad=.12); save(fig,path)


def fig3(path):
    # 所有框体增高，主箭头严格从框边缘起止，底部闭环独立于档案分流。
    fig,ax=plt.subplots(figsize=(6.5,4.55)); ax.set_xlim(0,12); ax.set_ylim(0,8); ax.axis("off")
    c="#3D4C5C"
    paper_box(ax,.32,5.12,2.50,.95,"工程精算候选",13.0)
    paper_box(ax,3.55,6.10,2.65,1.00,"严格档案 A_f\nV=0",12.7)
    paper_box(ax,3.55,3.62,2.65,1.00,"近可行档案 A_n\nV>0",12.7)
    paper_box(ax,7.02,5.12,2.15,.95,"父代选择",13.0)
    paper_box(ax,9.90,5.12,1.75,.95,"生成候选",12.5)
    arrow(ax,(2.82,5.78),(3.55,6.60),c,lw=1.1); arrow(ax,(2.82,5.42),(3.55,4.12),c,lw=1.1)
    arrow(ax,(6.20,6.60),(7.02,5.78),c,lw=1.1); arrow(ax,(6.20,4.12),(7.02,5.42),c,lw=1.1)
    arrow(ax,(9.17,5.60),(9.90,5.60),c,lw=1.1)
    arrow(ax,(10.78,5.12),(10.78,1.48),c,lw=1.1); arrow(ax,(10.78,1.48),(1.55,1.48),c,lw=1.1); arrow(ax,(1.55,1.48),(1.55,5.12),c,lw=1.1)
    fig.tight_layout(pad=.14); save(fig,path)


def fig4(path):
    scenes=["A","B","C","D","E"]
    bga=[16.7,60,100,40,3.3]; mga=[20,80,100,100,100]
    bga_n=["5/30","18/30","30/30","12/30","1/30"]; mga_n=["6/30","24/30","30/30","30/30","30/30"]
    bga_c=[2754.5,7066.5,10000,8000,8000]; mga_c=[663,2220.5,2074,3580.5,4314]
    fig,(a,b)=plt.subplots(1,2,figsize=(6.35,2.48),gridspec_kw={"wspace":.46}); x=np.arange(5); w=.33
    for ax,l,r,ylabel,ymax in ((a,bga,mga,"严格可行率（%）",108),(b,bga_c,mga_c,"调用中位数（次）",10800)):
        ax.bar(x-w/2,l,w,color="#AAB9C7",label="BGA"); ax.bar(x+w/2,r,w,color=TEAL,label="MGA")
        ax.set_xticks(x,scenes,fontsize=13); ax.set_ylim(0,ymax); ax.set_ylabel(ylabel,fontsize=12.5); ax.tick_params(axis="y",labelsize=10.5)
        ax.grid(axis="y",color=GRID,lw=.6); ax.spines[["top","right"]].set_visible(False)
    # 运行次数已在表 6 报告；图中只保留可读的主指标，避免柱顶数字在单栏中重叠。
    a.legend(frameon=False,fontsize=9.5,ncol=2,loc="upper left"); b.legend(frameon=False,fontsize=9.5,ncol=2,loc="upper right")
    a.set_title("（a）严格可行发现率",loc="left",fontsize=11.5,fontweight="bold",color=INK); b.set_title("（b）后续精算调用（描述性）",loc="left",fontsize=11.5,fontweight="bold",color=INK)
    fig.tight_layout(pad=.28); save(fig,path)


def fig5(path):
    rows=[("B  (n=14)",0,-2.83,1.27,"0.00%"),("C  (n=30)",-0.83,-1.03,-0.83,"-0.83%"),("D  (n=12)",-7.02,-11.96,-3.75,"-7.02%")]
    fig,ax=plt.subplots(figsize=(6.35,3.75)); ax.axvline(0,color="#73808C",lw=1.0); ys=np.array([2,1,0])
    for y,(name,est,lo,hi,label) in zip(ys,rows):
        ax.hlines(y,lo,hi,color="#304F69",lw=2.0); ax.vlines([lo,hi],y-.075,y+.075,color="#304F69",lw=1.15)
        ax.scatter(est,y,s=32,color="#304F69",zorder=3); ax.text(2.35,y,label,va="center",ha="right",fontsize=10.5,color="#263746")
    ax.set_xlim(-13,2.55); ax.set_ylim(-.45,2.45); ax.set_yticks(ys,labels=[r[0] for r in rows],fontsize=12)
    ax.set_xticks([-12,-9,-6,-3,0]); ax.tick_params(axis="x",labelsize=11)
    ax.set_xlabel("Hodges–Lehmann 成本变化率（%）",fontsize=12)
    ax.spines[["top","right"]].set_visible(False); ax.grid(axis="x",color="#E1E6EB",lw=.65); fig.tight_layout(pad=.32); save(fig,path)


def fig6(path):
    labels=["S0","S1","S2","S3","S4"]
    data={"A":[16.7,20,20,43.3,63.3],"B":[60,63.3,80,96.7,70],"D":[40,100,100,100,100]}; colors={"A":ORANGE,"B":BLUE,"D":TEAL}; marks={"A":"o","B":"s","D":"^"}
    # 提高图高并把图例置于绘图区上方，避免场景说明压到 S0--S4 数据点。
    fig,ax=plt.subplots(figsize=(6.35,3.55)); x=np.arange(5)
    for name,vals in data.items(): ax.scatter(x,vals,s=38,marker=marks[name],color=colors[name],label=f"场景 {name}",zorder=3)
    ax.set_xticks(x,labels,fontsize=12); ax.set_ylim(-4,112); ax.set_yticks([0,50,100]); ax.set_ylabel("严格可行率（%）",fontsize=12); ax.tick_params(axis="y",labelsize=10.5)
    ax.grid(axis="y",color=GRID,lw=.65); ax.spines[["top","right"]].set_visible(False)
    ax.legend(frameon=False,fontsize=10.0,ncol=3,loc="lower left",bbox_to_anchor=(0,1.03),borderaxespad=0,handletextpad=.32,columnspacing=1.0)
    fig.subplots_adjust(left=.13,right=.98,bottom=.22,top=.78); save(fig,path)


def fig7(path):
    first=json.loads((DATA/"03_主比较_BGA_MGA"/"diversity_failure_summary.json").read_text(encoding="utf-8")); second=json.loads((DATA/"03_主比较_BGA_MGA"/"diversity_failure_runs_t2rad_runs_t2r429.json").read_text(encoding="utf-8")); rows=first["diversity"]+second["diversity"]
    ids=["S-429","S-387","S-409","S-469","S-487"]; names=list("ABCDE"); fig,axs=plt.subplots(1,2,figsize=(6.35,2.75),gridspec_kw={"wspace":.52})
    for ax,field,title,ylabel,ymax in ((axs[0],"final_distance","（a）工作种群平均汉明距离","平均汉明距离 d_H",4.2),(axs[1],"signature_median","（b）不重复完整记录签名数","不重复完整记录签名数",15)):
        rel=[r for r in rows if field in r]; x=np.arange(5); bg=[next((r[field] for r in rel if r["scene"]==s and r["method"]=="BGA"),np.nan) for s in ids]; mg=[next((r[field] for r in rel if r["scene"]==s and r["method"]=="MGA"),np.nan) for s in ids]
        ax.scatter(x-.12,bg,s=48,color="#7B8FA4",label="BGA",zorder=3); ax.scatter(x+.12,mg,s=48,color=TEAL,label="MGA",zorder=3)
        ax.set_xticks(x,names,fontsize=13); ax.set_ylim(0,ymax); ax.set_ylabel(ylabel,fontsize=13); ax.tick_params(axis="y",labelsize=11.5); ax.set_title(title,fontsize=13.5,loc="left",fontweight="bold"); ax.grid(axis="y",color=GRID,lw=.6); ax.spines[["top","right"]].set_visible(False)
    axs[0].legend(frameon=False,fontsize=11.5,ncol=2,loc="upper left"); fig.tight_layout(pad=.28); save(fig,path)


def fig8(path):
    rows=[("A-BGA\n(n=25)",{"负载损耗":23,"空载损耗":1,"其他":1}),("A-MGA\n(n=24)",{"负载损耗":24}),("B-BGA\n(n=12)",{"空载损耗":6,"阻抗":3,"负载损耗":2,"其他":1}),("B-MGA\n(n=6)",{"负载损耗":5,"空载损耗":1}),("D-BGA\n(n=18)",{"空载损耗":7,"散热器中心距":5,"温升/电流":5,"阻抗":1}),("E-BGA\n(n=29)",{"散热器中心距":10,"温升/电流":11,"负载损耗":5,"空载损耗":1,"其他":2})]
    cats=["负载损耗","空载损耗","阻抗","散热器中心距","温升/电流","其他"]; cols={"负载损耗":BLUE,"空载损耗":TEAL,"阻抗":ORANGE,"散热器中心距":"#7C3AED","温升/电流":"#C94343","其他":"#6B7280"}
    fig,ax=plt.subplots(figsize=(6.35,2.45)); y=np.arange(len(rows)); left=np.zeros(len(rows))
    for cat in cats:
        vals=np.array([r[1].get(cat,0)/sum(r[1].values())*100 for r in rows]); ax.barh(y,vals,left=left,height=.6,color=cols[cat],label=cat); left+=vals
    ylabels=["A·BGA (25)","A·MGA (24)","B·BGA (12)","B·MGA (6)","D·BGA (18)","E·BGA (29)"]
    ax.set_yticks(y,ylabels,fontsize=14); ax.invert_yaxis(); ax.set_xlim(0,100); ax.set_xlabel("失败运行中主导约束构成（%）",fontsize=13); ax.tick_params(axis="x",labelsize=11.5); ax.grid(axis="x",color=GRID,lw=.6); ax.spines[["top","right"]].set_visible(False)
    ax.legend(frameon=False,fontsize=12.4,ncol=3,loc="lower center",bbox_to_anchor=(.5,1.02),columnspacing=.8,handlelength=1.25)
    fig.tight_layout(pad=.18); save(fig,path)


def figA1(path):
    fig,ax=plt.subplots(figsize=(6.25,2.05)); ax.set_xlim(0,12); ax.set_ylim(0,4); ax.axis("off")
    box(ax,.25,1.35,2.45,1.35,"13 条记录","可逐字段重放",PALE_BLUE,BLUE,12,8.5)
    box(ax,3.25,1.35,2.45,1.35,"32 字段 / 条","冻结输入可比",PALE_BLUE,BLUE,12,8.5)
    box(ax,6.25,1.35,2.45,1.35,"416 / 416","Java—Python 一致",PALE_TEAL,TEAL,12,8.5)
    box(ax,9.25,1.35,2.45,1.35,"0 项","字段不一致",PALE_TEAL,TEAL,12,8.5)
    arrow(ax,(2.7,2.02),(3.25,2.02)); arrow(ax,(5.7,2.02),(6.25,2.02)); arrow(ax,(8.7,2.02),(9.25,2.02))
    fig.tight_layout(pad=.12); save(fig,path)


def replace_blob(doc,index,path):
    p=doc.paragraphs[index]; ids=p._p.xpath('.//a:blip/@r:embed')
    if len(ids)!=1: raise RuntimeError((index,ids))
    doc.part.related_parts[ids[0]]._blob=path.read_bytes()


def clear(p):
    for child in list(p._p):
        if child.tag != qn("w:pPr"): p._p.remove(child)


def set_run(run,size=9):
    run.font.size=Pt(size); run.font.name="Times New Roman"; rpr=run._element.get_or_add_rPr(); fonts=rpr.rFonts
    if fonts is None: fonts=OxmlElement("w:rFonts"); rpr.insert(0,fonts)
    for attr,val in (("ascii","Times New Roman"),("hAnsi","Times New Roman"),("eastAsia","宋体"),("cs","Times New Roman")): fonts.set(qn(f"w:{attr}"),val)


def replace_text(p,value,size=9):
    clear(p); r=p.add_run(value); set_run(r,size)


def main():
    ASSET.mkdir(parents=True,exist_ok=True)
    renderers={23:("图1_工程优化闭环.png",fig1),28:("图2_完整记录整体继承.png",fig2),42:("图3_双档案流程.png",fig3),67:("图4_主比较结果.png",fig4),71:("图5_HL成本变化率.png",fig5),77:("图6_策略变体点图.png",fig6),79:("图7_多样性点图.png",fig7),84:("图8_失败约束构成.png",fig8),103:("图A1_计算链一致性.png",figA1)}
    paths={}
    for index,(name,fn) in renderers.items():
        paths[index]=ASSET/name; fn(paths[index])
    doc=Document(SOURCE)
    for index,path in paths.items(): replace_blob(doc,index,path)
    replace_text(doc.paragraphs[4],"针对三相油浸式变压器离散设计中材料目录离散、结构参数条件依赖及工程精算代价高的问题，提出一种目录约束双档案遗传优化方法。该方法以完整目录记录为原子决策单元，并通过合法化函数派生层数、油道和冷却关联变量，避免遗传操作产生无法由业务目录还原的虚拟规格。严格可行档案仅保存满足全部硬约束的候选，近可行档案保留可计算但存在约束超限的边界候选；覆盖初始化、耦合交叉和运行内已评价候选池注入构成默认增强搜索策略，诊断反馈局部变异作为可选扩展策略单独评估。主比较评价MGA默认增强策略整体，不将任何差异单独归因于双档案。基于五类预定义工程场景、每个场景30个配对随机种子的测试结果表明：相较目录约束基础遗传算法（BGA），在大容量长圆、散热器工况甲和工况乙中，MGA默认增强策略与BGA的严格可行发现率差异经双侧McNemar精确检验，并在5场景检验族内采用Holm校正后，p值分别为3.05e-5和1.86e-8；在双方共同严格成功的配对中，中容量长圆、波纹场景和大容量长圆、散热器工况甲的最低严格成本Hodges-Lehmann变化率分别为-0.83%和-7.02%。中小容量长圆、波纹场景的成本差异未达到统计显著。五个主场景中，BGA与MGA的第0代后真实精算调用中位数之比为1.85至4.82。结果支持：在当前冻结目录、价格快照与预算上限下，MGA默认增强策略可提高部分复杂配置的严格可行方案发现率；后续真实精算调用量较低仅为当前停止规则下的观测结果，不能替代首次严格解发现效率的判断。结论不外推为全部产品配置上的统一性能保证。")
    replace_text(doc.paragraphs[7],"A catalogue-constrained dual-archive genetic optimization method is proposed for three-phase oil-immersed transformer design. Complete catalogue records are atomic decision units, and legalization derives layer, duct, and cooling variables to prevent virtual specifications that cannot be restored from business catalogues. Strict-feasible and near-feasible archives retain hard-constraint-satisfying and computable boundary candidates, respectively. Coverage initialization, coupled crossover, and run-time evaluated-pool injection form the default strategy; diagnostic local mutation is an optional extension. The primary comparison evaluates the overall MGA-default strategy and does not isolate the independent effect of the dual archive. Five predefined engineering scenarios, each with 30 paired seeds, compare it with a catalogue-constrained baseline GA (BGA). In two large-capacity long-oval radiator scenarios, strict-feasible discovery-rate differences were tested by two-sided exact McNemar tests with five-scenario Holm adjustment (p=3.05e-5 and 1.86e-8). For jointly strict-feasible paired runs, Hodges-Lehmann cost-change estimates are -0.83% in the medium-capacity long-oval corrugated-tank scenario and -7.02% in large-capacity long-oval radiator scenario I; the small-to-medium-capacity long-oval corrugated-tank result is not significant. The ratio of median post-initialization engineering evaluations for BGA to MGA ranges from 1.85 to 4.82 across five scenarios. Under frozen catalogue and price snapshots and the same maximum actual-evaluation budget, the results support improved strict-feasible discovery for some of the tested complex configurations. Evaluation counts are descriptive under the stopping rule and do not measure first strict-feasible discovery efficiency; conclusions do not guarantee uniform performance.")
    replace_text(doc.paragraphs[11],"本文主比较结果来源于正式冻结后的运行轨迹；主实验冻结后独立开展Java与Python一致性核验及泛化边界检查，二者仅用于界定计算链一致性与适用边界，不参与主比较统计。主比较覆盖五类来自实际业务系统的工程配置快照：场景A为50 kVA圆形铁芯、波纹油箱，场景B为100 kVA长圆铁芯、波纹油箱，场景C为300 kVA长圆铁芯、波纹油箱，场景D为2 000 kVA长圆铁芯、散热器油箱工况甲，场景E为2 000 kVA长圆铁芯、散热器油箱工况乙。场景D与场景E的容量和主结构相同，但冻结的页面条件、目录快照与实际搜索规模不同，故作为独立场景分别报告。每个场景在BGA与MGA下采用相同的30个随机种子配对运行。另设一个200 kVA圆形铁芯、波纹油箱边界场景，双方均未得到严格解，仅用于边界检查，不进入主比较和显著性结论。")
    replace_text(doc.paragraphs[13],doc.paragraphs[13].text.replace("5个真实工程配置","5个实际业务系统冻结配置"))
    replace_text(doc.paragraphs[58],doc.paragraphs[58].text.replace("真实工程配置快照","实际业务系统冻结配置快照"))
    replace_text(doc.paragraphs[73],"表6 BGA与MGA严格成功情况及共同成功条件下的成本比较",9)
    replace_text(doc.paragraphs[78],"图6 独立扩展批次中预定义策略配置的严格可行率（描述性结果；S0：基础BGA；S1：双档案+耦合交叉；S2：S1+候选池注入；S3：S2+诊断探测；S4：S3关闭独立近可行档案。各点为独立策略配置，不表示连续递进关系；仅展示场景A、B、D，场景C为100.0%，场景E扩展为0.0%，完整结果见表8）",9)
    replace_text(doc.paragraphs[79],"图7 五个主场景的最终工作种群平均汉明距离与严格档案不重复完整记录签名数（描述性统计）",9)
    replace_text(doc.paragraphs[93],doc.paragraphs[93].text.replace("当前真实测试支持","当前主比较测试支持"))
    replace_text(doc.paragraphs[95],doc.paragraphs[95].text.replace("5个真实工程配置","5个实际业务系统冻结配置"))
    replace_text(doc.paragraphs[101],doc.paragraphs[101].text.replace("在当前可逐位复现的13条三相历史记录范围内，","从当前能够在Java与Python两套冻结计算链中逐字段重放的三相历史记录中选取13条，"))
    replace_text(doc.paragraphs[105],"附录B 五个实际业务系统配置场景的冻结搜索快照",10.5)
    replace_text(doc.paragraphs[106],doc.paragraphs[106].text.replace("五个真实主比较场景","五个实际业务系统主比较场景").replace("当前归档尚未建立独立snapshot UUID，以“归档日期+场景+对应SQLite轨迹/JSON汇总文件”共同识别一次冻结运行快照。","当前归档尚未建立独立snapshot UUID，以“归档日期+场景+对应SQLite轨迹/JSON汇总文件”共同识别一次冻结运行快照。企业目录与价格快照不对外公开，本文的可复核性限定为内部冻结归档；外部复现需基于脱敏配置摘要及授权元数据。"))
    replace_text(doc.paragraphs[85],"图8 未获得严格可行解运行的主导约束构成（描述性统计；分母为各方法未获严格解的运行数，具体计数见表9；场景D/E的MGA无失败运行，未绘制百分比条）",9)
    doc.save(OUTPUT); print(OUTPUT)

if __name__=="__main__": main()

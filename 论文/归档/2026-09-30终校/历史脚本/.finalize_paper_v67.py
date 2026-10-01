"""V6.7：以 Word OMath 统一核心变量，并强化图4单栏可读性。"""
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = Path(r"D:\IdeaProject\faladi\three_phase_ga_optimizer")
SOURCE = ROOT / "论文" / "目录约束双档案遗传优化三相变压器_版式与引用优化版V6.6-送外审微调终校版.docx"
OUTPUT = ROOT / "论文" / "目录约束双档案遗传优化三相变压器_版式与引用优化版V6.7-送外审数学对象与图4终校版.docx"
ASSET = ROOT / "论文" / "图表-送外审终校V6.7"

INK, TEAL, GRID = "#19324D", "#167C74", "#D9E3EC"
plt.rcParams.update({"font.sans-serif": ["Microsoft YaHei", "SimHei", "Arial Unicode MS"], "axes.unicode_minus": False})


def set_text_font(run, size=9):
    run.font.size = Pt(size); run.font.name = "Times New Roman"
    rpr = run._element.get_or_add_rPr(); fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts"); rpr.insert(0, fonts)
    fonts.set(qn("w:ascii"), "Times New Roman"); fonts.set(qn("w:hAnsi"), "Times New Roman")
    fonts.set(qn("w:eastAsia"), "宋体"); fonts.set(qn("w:cs"), "Times New Roman")
    rpr.set(qn("w:hint"), "eastAsia")


def clear(p):
    for child in list(p._p):
        if child.tag != qn("w:pPr"):
            p._p.remove(child)


def text(p, value, size=9):
    r = p.add_run(value); set_text_font(r, size); return r


def m_run(value):
    r = OxmlElement("m:r")
    rpr = OxmlElement("m:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), "Cambria Math"); rfonts.set(qn("w:hAnsi"), "Cambria Math")
    rpr.append(rfonts); r.append(rpr)
    t = OxmlElement("m:t"); t.text = value; r.append(t)
    return r


def math_plain(p, value):
    o = OxmlElement("m:oMath"); o.append(m_run(value)); p._p.append(o)


def math_sub(p, base, sub, suffix=""):
    o = OxmlElement("m:oMath"); s = OxmlElement("m:sSub")
    e = OxmlElement("m:e"); e.append(m_run(base))
    low = OxmlElement("m:sub"); low.append(m_run(sub))
    s.append(e); s.append(low); o.append(s)
    if suffix:
        o.append(m_run(suffix))
    p._p.append(o)


def body_setup(p):
    clear(p); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    f=p.paragraph_format; f.first_line_indent=Pt(18); f.line_spacing=1.15; f.space_before=Pt(0); f.space_after=Pt(0)


def replace_para35(doc):
    p=doc.paragraphs[35]; body_setup(p)
    text(p,"严格可行档案记为"); math_sub(p,"A","f")
    text(p,"，其中仅保留calculable=True、complete=True且全部工程约束均满足的候选，并按总成本从低到高排序。近可行档案记为"); math_sub(p,"A","n")
    text(p,"，其中保留calculable=True、complete=True、但至少一项约束超限的候选。设"); math_plain(p,"I(q)")
    text(p,"为当前页面配置实际提供上下限的约束集合；配置未提供范围时不虚构该项硬约束。对"); math_sub(p,"A","n")
    text(p,"内候选按"); math_plain(p,"κ(x)=(V(x),")
    math_sub(p,"V","max","(x),"); math_sub(p,"n","v","(x),C(x))")
    text(p,"进行升序字典序排序：首先比较归一化超限总和"); math_plain(p,"V(x)")
    text(p,"，若相同则依次比较最大单项超限"); math_sub(p,"V","max","(x)")
    text(p,"、违规项数量"); math_sub(p,"n","v","(x)")
    text(p,"和总成本"); math_plain(p,"C(x)")
    text(p,"。两档案容量均固定为"); math_plain(p,"L=24")
    text(p,"，作为独立于工作种群规模的截断参数；工作种群规模则随第0代候选池规模设定。"); math_sub(p,"A","f")
    text(p,"是优化器最终输出候选的唯一来源，"); math_sub(p,"A","n")
    text(p,"提供跨越约束边界的父代机会；仅在诊断扩展策略启用时，还可作为诊断反馈输入。")


def replace_para52(doc):
    p=doc.paragraphs[52]; body_setup(p)
    text(p,"运行内已评价候选池是指本次运行中已完成工程精算且完整的候选集合，不是外部训练集，也不是目标配置既有结果。每一代注入时，算法仅从该候选池中随机抽取基体和供体；候选池注入的单模块替换概率")
    math_sub(p,"p","single","=0.75")
    text(p,"，即整体替换结构和电磁耦合块或冷却块之一，双模块替换概率")
    math_sub(p,"p","double","=0.25")
    text(p,"，即同时替换这两个模块，再经合法化函数合法化后形成未见候选。")
    math_sub(p,"p","single")
    text(p,"与表5中的耦合交叉率")
    math_sub(p,"p","cross","=0.75")
    text(p,"为不同的冻结参数：前者限定候选池注入时替换模块的数量，后者限定遗传交叉采用耦合方式的概率。若该候选池少于两条，才回退到原始页面允许域的完整合法随机候选。历史结构种子只用于第0代：目标配置ID被排除，候选配置须具有相同产品类型和频率；其余特征以容量、额定电压、绕法、联结方式、铁芯形状/直径、油箱类型和高低压导线类别等加权评分，历史记录还必须能映射为当前目录记录、通过当前搜索域校验，并由当前精算器重新评价。因此，历史结构种子与运行内已评价候选池相互独立；历史价格和合规等级仅用于决定种子的尝试顺序，不进入当前目标函数、严格可行判定或当前场景的工程性能标签，因而不向目标配置引入结果标签信息。")


def replace_para58_70(doc):
    # 只将正文核心统计变量替换为 OMath；其余文字和结论不动。
    for index, source, specs in (
        (58, doc.paragraphs[58].text, [("N0", "N", "0"), ("B′", "B", None)]),
        (70, doc.paragraphs[70].text, [("r_rb", "r", "rb")]),
    ):
        p=doc.paragraphs[index]; body_setup(p)
        cursor=0
        tokens=sorted(specs, key=lambda x:len(x[0]), reverse=True)
        while cursor<len(source):
            matched=None
            for token,base,sub in tokens:
                if source.startswith(token,cursor): matched=(token,base,sub); break
            if not matched:
                end=cursor+1
                while end<len(source) and not any(source.startswith(t[0],end) for t in tokens): end+=1
                text(p,source[cursor:end]); cursor=end; continue
            token,base,sub=matched
            if sub is None: math_plain(p,base+"′")
            else: math_sub(p,base,sub)
            cursor+=len(token)


def render_fig4(path):
    scenes=["A","B","C","D","E"]
    bga_success=[16.7,60.0,100.0,40.0,3.3]; mga_success=[20.0,80.0,100.0,100.0,100.0]
    bga_calls=[2754.5,7066.5,10000.0,8000.0,8000.0]; mga_calls=[663.0,2220.5,2074.0,3580.5,4314.0]
    fig,(a1,a2)=plt.subplots(1,2,figsize=(6.4,2.45),gridspec_kw={"wspace":.48})
    x=np.arange(5); width=.35
    for ax,left,right,ylabel,ymax in ((a1,bga_success,mga_success,"严格可行率（%）",110),(a2,bga_calls,mga_calls,"后续真实精算调用中位数（次）",11000)):
        ax.bar(x-width/2,left,width,label="BGA",color="#AAB9C7")
        ax.bar(x+width/2,right,width,label="MGA",color=TEAL)
        ax.set_xticks(x,scenes,fontsize=20); ax.set_ylabel(ylabel,fontsize=19); ax.set_ylim(0,ymax)
        ax.tick_params(axis="y",labelsize=17); ax.grid(axis="y",color=GRID,linewidth=.6); ax.spines[["top","right"]].set_visible(False)
    a1.legend(frameon=False,fontsize=17,ncol=2,loc="upper left"); a2.legend(frameon=False,fontsize=17,ncol=2,loc="upper right")
    fig.tight_layout(pad=.25); fig.savefig(path,dpi=440,bbox_inches="tight",facecolor="white"); plt.close(fig)


def replace_image(doc, paragraph_index, path):
    p=doc.paragraphs[paragraph_index]; ids=p._p.xpath(".//a:blip/@r:embed")
    if len(ids)!=1: raise RuntimeError(f"image paragraph {paragraph_index}: {ids}")
    doc.part.related_parts[ids[0]]._blob=path.read_bytes()


def main():
    ASSET.mkdir(parents=True,exist_ok=True)
    fig4=ASSET/"图4_单栏大字版.png"; render_fig4(fig4)
    doc=Document(SOURCE)
    replace_para35(doc); replace_para52(doc); replace_para58_70(doc)
    replace_image(doc,67,fig4)
    doc.save(OUTPUT); print(OUTPUT)

if __name__=="__main__": main()

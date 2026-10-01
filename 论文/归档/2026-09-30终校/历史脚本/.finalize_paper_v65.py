"""V6.5：送外审前最后的符号、关键词与参考文献页校正。

基于 V6.4，仅修正表 1 的符号—含义映射、数学变量格式、表 2 的默认/扩展边界、
首页关键词和参考文献末页双栏平衡；不修改实验数据、统计值或正文结论。
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(r"D:\IdeaProject\faladi\three_phase_ga_optimizer")
SOURCE = ROOT / "论文" / "目录约束双档案遗传优化三相变压器_版式与引用优化版V6.4-送外审版式与图表可读性终校版.docx"
OUTPUT = ROOT / "论文" / "目录约束双档案遗传优化三相变压器_版式与引用优化版V6.5-送外审符号与末页终校版.docx"
BODY = 9.0
TABLE = 7.5


def set_fonts(run, size, east_asia="宋体", latin="Times New Roman", math=False, bold=None):
    run.font.size = Pt(size)
    run.font.name = "Cambria Math" if math else latin
    if bold is not None:
        run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    name = "Cambria Math" if math else latin
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name if math else east_asia)
    rfonts.set(qn("w:cs"), name)
    rpr.set(qn("w:hint"), "default" if math else "eastAsia")


def clear_paragraph(p):
    for child in list(p._p):
        if child.tag != qn("w:pPr"):
            p._p.remove(child)


def add_text(p, text, size=BODY, bold=None):
    r = p.add_run(text)
    set_fonts(r, size, bold=bold)
    return r


def add_math(p, base, sub=None, suffix="", size=BODY):
    r = p.add_run(base)
    set_fonts(r, size, math=True)
    if sub:
        s = p.add_run(sub)
        set_fonts(s, size, math=True)
        s.font.subscript = True
    if suffix:
        r = p.add_run(suffix)
        set_fonts(r, size, math=True)


def rebuild_with_tokens(p, replacements, size=BODY, indent=True):
    """按精确 token 重建段落；replacements 值为 (base, sub, suffix) 或 None。"""
    original = p.text
    tokens = sorted(replacements, key=len, reverse=True)
    parts = []
    cursor = 0
    while cursor < len(original):
        hit = None
        for token in tokens:
            if original.startswith(token, cursor):
                hit = token
                break
        if hit is None:
            end = cursor + 1
            while end < len(original) and not any(original.startswith(t, end) for t in tokens):
                end += 1
            parts.append(("text", original[cursor:end]))
            cursor = end
        else:
            parts.append(("math", replacements[hit]))
            cursor += len(hit)
    alignment = p.alignment if p.alignment is not None else WD_ALIGN_PARAGRAPH.JUSTIFY
    clear_paragraph(p)
    p.alignment = alignment
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.first_line_indent = Pt(18) if indent else Pt(0)
    for kind, value in parts:
        if kind == "text":
            add_text(p, value, size=size)
        else:
            add_math(p, value[0], value[1], value[2], size=size)


def set_cell_symbol(cell, base, sub=None, suffix=""):
    p = cell.paragraphs[0]
    clear_paragraph(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    add_math(p, base, sub, suffix, size=TABLE)


def set_cell_text(cell, text):
    p = cell.paragraphs[0]
    clear_paragraph(p)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    add_text(p, text, size=TABLE)


def fix_symbol_table(doc):
    t = doc.tables[0]
    # 行/列采用三组“符号—含义”布局；逐项重建，避免此前符号列错位。
    entries = [
        ((1, 0, "q", None, ""), (1, 1, "页面配置与固定工艺条件")),
        ((1, 2, "D", None, ""), (1, 3, "材料、结构与冷却目录快照")),
        ((1, 4, "p", None, ""), (1, 5, "材料价格与计算参数快照")),
        ((2, 0, "z", None, ""), (2, 1, "由Φ派生的层数、油道等变量向量")),
        ((2, 2, "Φ", None, ""), (2, 3, "候选合法化函数")),
        ((2, 4, "Ω", None, ""), (2, 5, "满足页面与目录约束的实际候选域")),
        ((3, 0, "E", None, ""), (3, 1, "单设备工程精算器")),
        ((3, 2, "g", "i", " / h"), (3, 3, "第i个不等式 / 第j个等式工程约束")),
        ((3, 4, "κ", None, ""), (3, 5, "近可行档案字典序排序键")),
        ((4, 0, "d̄", "H", ""), (4, 1, "四模块平均汉明距离（0—4）")),
        ((4, 2, "H", "m", ""), (4, 3, "模块取值熵")),
        ((4, 4, "B", None, "′"), (4, 5, "第0代后真实精算调用总预算")),
    ]
    for symbol, meaning in entries:
        r, c, base, sub, suffix = symbol
        set_cell_symbol(t.cell(r, c), base, sub, suffix)
        r, c, text = meaning
        set_cell_text(t.cell(r, c), text)
    # g_i / h_j 需要两段独立下标，不能用普通字符串“g_i / h_j”替代。
    p = t.cell(3, 2).paragraphs[0]
    clear_paragraph(p); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0); p.paragraph_format.line_spacing = 1.0
    add_math(p, "g", "i", size=TABLE)
    add_text(p, " / ", size=TABLE)
    add_math(p, "h", "j", size=TABLE)


def fix_keywords(doc):
    p = doc.paragraphs[5]
    clear_paragraph(p)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    add_text(p, "关键词：", size=BODY, bold=True)
    add_text(p, "三相油浸式变压器；目录约束；约束优化；离散优化；遗传算法；双档案；工程精算一致性", size=BODY)


def fix_method_terms(doc):
    # 表 2：默认 MGA 不含诊断探测，明确“可选”边界。
    set_cell_text(doc.tables[1].cell(4, 2), "双档案、候选池与可选诊断反馈")

    # 相关工作与方法/结果中的核心变量恢复数学上下标与希腊字母。
    rebuild_with_tokens(doc.paragraphs[31], {"epsilon": ("ε", None, "")})
    rebuild_with_tokens(doc.paragraphs[35], {
        "A_f": ("A", "f", ""), "A_n": ("A", "n", ""), "I(q)": ("I(q)", None, ""),
        "V_max(x)": ("V", "max", "(x)"), "n_v(x)": ("n", "v", "(x)"),
        "V(x)": ("V(x)", None, ""), "L=24": ("L=24", None, ""),
        "排序键(x)": ("κ(x)", None, ""),
    })
    rebuild_normalized_violation(doc.paragraphs[37])
    rebuild_with_tokens(doc.paragraphs[52], {
        "p_single": ("p", "single", ""), "p_double": ("p", "double", ""), "p_cross": ("p", "cross", ""),
    })
    rebuild_with_tokens(doc.paragraphs[58], {
        "N0": ("N", "0", ""), "B'": ("B", None, "′"),
    })
    rebuild_with_tokens(doc.paragraphs[70], {
        "r_rb": ("r", "rb", ""),
    })


def rebuild_normalized_violation(p):
    """重建 2.1.1 的长段，保留原文字并将求和/下标写成独立数学运行。"""
    source = p.text
    marker = "对完整且可计算的候选，"
    formula_end = "并据此进入A_f或A_n排序。"
    before, remainder = source.split(marker, 1)
    _, after = remainder.split(formula_end, 1)
    clear_paragraph(p)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.first_line_indent = Pt(18)
    # 公式前的所有归一化解释不改写，只把常用变量以内联数学格式呈现。
    prefix = before + marker
    for text, math in ((prefix, False), ("V(x)=", True)):
        (add_math if math else add_text)(p, text, size=BODY)
    add_math(p, "∑", "i∈I(q)", size=BODY); add_math(p, "v", "i", "(x)", size=BODY)
    add_text(p, "，", size=BODY)
    add_math(p, "V", "max", "(x)=max", size=BODY); add_math(p, "i∈I(q)", size=BODY)
    add_math(p, "v", "i", "(x)", size=BODY)
    add_text(p, "，", size=BODY)
    add_math(p, "n", "v", "(x)=|{i∈I(q):", size=BODY); add_math(p, "v", "i", "(x)>0}|", size=BODY)
    add_text(p, "，并据此进入", size=BODY); add_math(p, "A", "f", size=BODY)
    add_text(p, "或", size=BODY); add_math(p, "A", "n", "", size=BODY)
    add_text(p, "排序。" + after, size=BODY)


def balance_final_reference_page(doc):
    """在最后一条参考文献前插入分栏符，让第10页左右栏均有内容。"""
    p = doc.paragraphs[129]  # [17]
    r = OxmlElement("w:r")
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "column")
    r.append(br)
    p._p.insert(1, r)


def main():
    doc = Document(SOURCE)
    fix_symbol_table(doc)
    fix_keywords(doc)
    fix_method_terms(doc)
    balance_final_reference_page(doc)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()

"""V6.9：送外审终校——补全表1数学对象，并提升图4单栏可读性。"""
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = Path(r"D:\IdeaProject\faladi\three_phase_ga_optimizer")
SOURCE = ROOT / "论文" / "目录约束双档案遗传优化三相变压器_版式与引用优化版V6.7-送外审数学对象与图4终校版.docx"
OUTPUT = ROOT / "论文" / "目录约束双档案遗传优化三相变压器_版式与引用优化版V6.9-送外审图4与数学变量终校版.docx"
ASSET = ROOT / "论文" / "图表-送外审终校V6.9"

INK, TEAL, GRID = "#19324D", "#167C74", "#D9E3EC"
plt.rcParams.update({"font.sans-serif": ["Microsoft YaHei", "SimHei", "Arial Unicode MS"], "axes.unicode_minus": False})


def set_font(run, size=7.5):
    run.font.size = Pt(size)
    run.font.name = "Times New Roman"
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr, value in (("ascii", "Times New Roman"), ("hAnsi", "Times New Roman"),
                        ("eastAsia", "宋体"), ("cs", "Times New Roman")):
        fonts.set(qn(f"w:{attr}"), value)
    rpr.set(qn("w:hint"), "eastAsia")


def clear_paragraph(p):
    for child in list(p._p):
        if child.tag != qn("w:pPr"):
            p._p.remove(child)


def add_text(p, value, size=7.5):
    run = p.add_run(value)
    set_font(run, size)


def m_run(value):
    run = OxmlElement("m:r")
    rpr = OxmlElement("m:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), "Cambria Math")
    rfonts.set(qn("w:hAnsi"), "Cambria Math")
    rpr.append(rfonts)
    run.append(rpr)
    node = OxmlElement("m:t")
    node.text = value
    run.append(node)
    return run


def add_math_plain(p, value):
    obj = OxmlElement("m:oMath")
    obj.append(m_run(value))
    p._p.append(obj)


def add_math_sub(p, base, sub, suffix=""):
    obj = OxmlElement("m:oMath")
    subscript = OxmlElement("m:sSub")
    base_node = OxmlElement("m:e")
    base_node.append(m_run(base))
    lower_node = OxmlElement("m:sub")
    lower_node.append(m_run(sub))
    subscript.append(base_node)
    subscript.append(lower_node)
    obj.append(subscript)
    if suffix:
        obj.append(m_run(suffix))
    p._p.append(obj)


def set_cell_math(cell, parts):
    p = cell.paragraphs[0]
    clear_paragraph(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for kind, values in parts:
        if kind == "text":
            add_text(p, values)
        elif kind == "plain":
            add_math_plain(p, values)
        elif kind == "sub":
            add_math_sub(p, *values)


def render_fig4(path):
    scenes = ["A", "B", "C", "D", "E"]
    bga_success = [16.7, 60.0, 100.0, 40.0, 3.3]
    mga_success = [20.0, 80.0, 100.0, 100.0, 100.0]
    bga_calls = [2754.5, 7066.5, 10000.0, 8000.0, 8000.0]
    mga_calls = [663.0, 2220.5, 2074.0, 3580.5, 4314.0]

    fig, (a1, a2) = plt.subplots(2, 1, figsize=(5.6, 3.85), gridspec_kw={"hspace": 0.78})
    x = np.arange(len(scenes))
    width = 0.34
    specs = (
        (a1, bga_success, mga_success, "严格可行率（%）", 110),
        (a2, bga_calls, mga_calls, "第0代后真实精算调用中位数（次）", 11000),
    )
    for index, (ax, left, right, ylabel, ymax) in enumerate(specs):
        ax.bar(x - width / 2, left, width, label="BGA", color="#AAB9C7")
        ax.bar(x + width / 2, right, width, label="MGA", color=TEAL)
        ax.set_xticks(x, scenes, fontsize=18)
        ax.set_ylabel(ylabel, fontsize=17, labelpad=7)
        ax.set_ylim(0, ymax)
        ax.tick_params(axis="y", labelsize=15)
        ax.grid(axis="y", color=GRID, linewidth=.7)
        ax.spines[["top", "right"]].set_visible(False)
        ax.legend(frameon=False, fontsize=15, ncol=2, loc="upper left")
        ax.text(-0.12, 1.05, "（a）" if index == 0 else "（b）", transform=ax.transAxes,
                fontsize=16, fontweight="bold", color=INK)
    fig.savefig(path, dpi=440, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def replace_image(doc, paragraph_index, image):
    paragraph = doc.paragraphs[paragraph_index]
    ids = paragraph._p.xpath(".//a:blip/@r:embed")
    if len(ids) != 1:
        raise RuntimeError(f"unexpected figure relationship at {paragraph_index}: {ids}")
    paragraph._p.xpath(".//wp:extent")[0].set("cx", "2834640")
    paragraph._p.xpath(".//wp:extent")[0].set("cy", "1554480")
    paragraph._p.xpath(".//a:xfrm/a:ext")[0].set("cx", "2834640")
    paragraph._p.xpath(".//a:xfrm/a:ext")[0].set("cy", "1554480")
    doc.part.related_parts[ids[0]]._blob = image.read_bytes()


def main():
    ASSET.mkdir(parents=True, exist_ok=True)
    image = ASSET / "图4_紧凑纵向双子图大字版.png"
    render_fig4(image)
    doc = Document(SOURCE)
    table = doc.tables[0]
    set_cell_math(table.cell(3, 2), [("sub", ("g", "i")), ("text", " / "), ("sub", ("h", "j"))])
    set_cell_math(table.cell(4, 0), [("sub", ("d̄", "H"))])
    set_cell_math(table.cell(4, 2), [("sub", ("H", "m"))])
    set_cell_math(table.cell(4, 4), [("plain", "B′")])
    replace_image(doc, 67, image)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()

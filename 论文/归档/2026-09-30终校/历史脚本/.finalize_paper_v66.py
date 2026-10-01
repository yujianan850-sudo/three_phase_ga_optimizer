"""V6.6：送外审前微调。

仅重输第2.4节首段以清除异常字距，并压缩表5的两行长参数说明；
核心变量均用 Cambria Math 下标运行，不改变实验数据、表格数值或结论。
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = Path(r"D:\IdeaProject\faladi\three_phase_ga_optimizer")
SOURCE = ROOT / "论文" / "目录约束双档案遗传优化三相变压器_版式与引用优化版V6.5-送外审符号与末页终校版.docx"
OUTPUT = ROOT / "论文" / "目录约束双档案遗传优化三相变压器_版式与引用优化版V6.6-送外审微调终校版.docx"


def set_fonts(run, size, math=False, bold=None):
    run.font.size = Pt(size)
    run.font.name = "Cambria Math" if math else "Times New Roman"
    if bold is not None:
        run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    family = "Cambria Math" if math else "Times New Roman"
    rfonts.set(qn("w:ascii"), family)
    rfonts.set(qn("w:hAnsi"), family)
    rfonts.set(qn("w:eastAsia"), family if math else "宋体")
    rfonts.set(qn("w:cs"), family)
    rpr.set(qn("w:hint"), "default" if math else "eastAsia")


def clear(p):
    for child in list(p._p):
        if child.tag != qn("w:pPr"):
            p._p.remove(child)


def text(p, value, size=9, bold=None):
    r = p.add_run(value)
    set_fonts(r, size, bold=bold)


def math(p, base, sub=None, suffix="", size=9):
    r = p.add_run(base)
    set_fonts(r, size, math=True)
    if sub is not None:
        s = p.add_run(sub)
        set_fonts(s, size, math=True)
        s.font.subscript = True
    if suffix:
        r = p.add_run(suffix)
        set_fonts(r, size, math=True)


def set_body(p, indent=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    clear(p)
    p.alignment = align
    fmt = p.paragraph_format
    fmt.first_line_indent = Pt(18) if indent else Pt(0)
    fmt.line_spacing = 1.15
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)


def rewrite_section_24(doc):
    """整段重输，消除历史 run 分裂和分散对齐导致的“MGA 默 认”字距。"""
    p = doc.paragraphs[55]
    set_body(p)
    text(p, "本文的主比较对象为目录约束基础遗传算法（BGA）与MGA默认增强策略（MGA-default，以下简称MGA）。二者共享候选域、合法化函数、单设备精算器、第0代覆盖策略、随机种子和第0代后真实精算总预算上限；主比较中的MGA额外启用独立近可行档案及其父代抽样、耦合交叉和运行内已评价候选池注入，变异模式为普通随机变异。诊断局部探测属于独立扩展策略批次，不是本节主比较MGA的组成部分。主比较用于检验该默认增强搜索机制相对BGA是否改善结果，不用于将任何差异单独归因于双档案。去双档案变体（single-archive variant）仅作为扩展对照：在诊断探测扩展策略的其余机制保持不变时关闭独立近可行档案；其余机制的leave-one-out对照尚待补充。")


def cell_parts(cell, parts):
    p = cell.paragraphs[0]
    clear(p)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    for kind, value in parts:
        if kind == "text":
            text(p, value, size=7.5)
        else:
            base, sub, suffix = value
            math(p, base, sub, suffix, size=7.5)


def streamline_table5(doc):
    t = doc.tables[6]
    # 用短分号段取代两句长说明，维持7.5pt并减少表格视觉密度。
    cell_parts(t.cell(6, 1), [
        ("text", "BGA：关闭"), ("math", ("A", "n", "")), ("text", "、"),
        ("math", ("p", "cross", "=0")), ("text", "，原始域随机注入、普通随机变异；MGA：启用"),
        ("math", ("A", "n", "")), ("text", "（父代B概率0.30）、"),
        ("math", ("p", "cross", "=0.75")), ("text", "与候选池注入（"),
        ("math", ("p", "inj", "=0.10；")), ("math", ("p", "single", "=0.75；")),
        ("math", ("p", "double", "=0.25")), ("text", "），普通随机变异。"),
    ])
    cell_parts(t.cell(7, 1), [
        ("text", "公共："), ("math", ("L", None, "=24")), ("text", "；末端细化为前3条严格精英、最多12邻居，计入"),
        ("math", ("B", None, "′")), ("text", "。默认MGA：random_only；诊断扩展：single_primary=0.65，每轮不大于3条。"),
    ])


def main():
    doc = Document(SOURCE)
    rewrite_section_24(doc)
    streamline_table5(doc)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()

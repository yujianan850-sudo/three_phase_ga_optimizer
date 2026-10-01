"""V6.11：修正正文居中段落、首行缩进及主比较定义句的历史格式碎片。"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = Path(r"D:\IdeaProject\faladi\three_phase_ga_optimizer")
SOURCE = ROOT / "论文" / "目录约束双档案遗传优化三相变压器_版式与引用优化版V6.10-送外审数学对象与图4终校版.docx"
OUTPUT = ROOT / "论文" / "目录约束双档案遗传优化三相变压器_版式与引用优化版V6.11-送外审正文对齐终校版.docx"


def set_font(run, size=9):
    run.font.size = Pt(size)
    run.font.name = "Times New Roman"
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr, value in (("ascii", "Times New Roman"), ("hAnsi", "Times New Roman"),
                        ("eastAsia", "宋体"), ("cs", "Times New Roman")):
        fonts.set(qn(f"w:{attr}"), value)
    rpr.set(qn("w:hint"), "eastAsia")


def set_body_layout(p):
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt = p.paragraph_format
    fmt.first_line_indent = Pt(18)
    fmt.line_spacing = 1.15
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)


def clear_paragraph(p):
    for child in list(p._p):
        if child.tag != qn("w:pPr"):
            p._p.remove(child)


def rewrite_main_comparison(p):
    value = (
        "本文主比较对象为目录约束基础遗传算法（BGA）和默认增强策略MGA-default（以下简称MGA）。"
        "二者共享候选域、合法化函数、单设备精算器、第0代覆盖策略、随机种子和第0代后真实精算总预算上限；"
        "主比较中的MGA额外启用独立近可行档案及其父代抽样、耦合交叉和运行内已评价候选池注入，变异模式为普通随机变异。"
        "诊断局部探测属于独立扩展策略批次，不是本节主比较MGA的组成部分。"
        "主比较用于检验该默认增强搜索机制相对BGA是否改善结果，不用于将任何差异单独归因于双档案。"
        "去双档案变体（single-archive variant）仅作为扩展对照：在诊断探测扩展策略的其余机制保持不变时关闭独立近可行档案；"
        "其余机制的leave-one-out对照尚待补充。"
    )
    clear_paragraph(p)
    set_body_layout(p)
    run = p.add_run(value)
    set_font(run)


def main():
    doc = Document(SOURCE)
    # 普通正文误设为居中的五段：统一为两端对齐、两字符首行缩进。
    for index in (28, 66, 76, 83, 106):
        set_body_layout(doc.paragraphs[index])
    # 重输这段，清除历史 run/特殊空格造成的“MGA 默认增强策略”字距异常。
    rewrite_main_comparison(doc.paragraphs[55])
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()

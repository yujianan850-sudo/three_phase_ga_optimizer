from __future__ import annotations

import shutil
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


PAPER_DIR = Path(r"D:\IdeaProject\faladi\three_phase_ga_optimizer\论文")
SOURCE = PAPER_DIR / "目录约束双档案遗传优化三相变压器_版式与引用优化版V5.9-投稿前终稿版式收紧.docx"
TARGET = PAPER_DIR / "目录约束双档案遗传优化三相变压器_版式与引用优化版V6.1-送审前字符与引用终校版.docx"


def set_fonts(run) -> None:
    """Use explicit CJK and Latin fonts to avoid viewer-dependent fallback glyphs."""
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), "Times New Roman")
    r_fonts.set(qn("w:hAnsi"), "Times New Roman")
    r_fonts.set(qn("w:eastAsia"), "宋体")
    r_fonts.set(qn("w:cs"), "Times New Roman")
    r_fonts.set(qn("w:hint"), "eastAsia")


def iter_paragraphs(parent):
    for paragraph in parent.paragraphs:
        yield paragraph
    for table in parent.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from iter_paragraphs(cell)


def replace_in_runs(paragraph, replacements: dict[str, str]) -> None:
    for run in paragraph.runs:
        text = run.text
        for old, new in replacements.items():
            text = text.replace(old, new)
        if text != run.text:
            run.text = text


def set_single_run_text(paragraph, text: str) -> None:
    # Only used on plain-text paragraphs (never paragraphs containing equations or drawings).
    for run in paragraph.runs:
        run._element.getparent().remove(run._element)
    run = paragraph.add_run(text)
    set_fonts(run)


def append_math_number(paragraph, number: int) -> None:
    """Keep the formula number in the same OMML paragraph as its formula."""
    omath = next((node for node in paragraph._p.iter() if node.tag == qn("m:oMath")), None)
    if omath is None:
        raise RuntimeError(f"formula ({number}) not found")
    math_run = OxmlElement("m:r")
    text = OxmlElement("m:t")
    text.set(qn("xml:space"), "preserve")
    text.text = f"     ({number})"
    math_run.append(text)
    omath.append(math_run)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def replace_math_text(paragraph, replacements: dict[str, str]) -> None:
    for node in paragraph._p.iter():
        if node.tag == qn("m:t") and node.text:
            value = node.text
            for old, new in replacements.items():
                value = value.replace(old, new)
            node.text = value


def remove_paragraph(paragraph) -> None:
    paragraph._element.getparent().remove(paragraph._element)


def set_cell_text(cell, text: str, size: float = 6.4) -> None:
    cell.text = text
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            set_fonts(run)
            run.font.size = Pt(size)


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    shutil.copy2(SOURCE, TARGET)
    document = Document(TARGET)

    # Explicitly select fonts for all ordinary text, table cells, and headers/footers.
    for paragraph in iter_paragraphs(document):
        for run in paragraph.runs:
            set_fonts(run)
    for section in document.sections:
        for story in (section.header, section.footer):
            for paragraph in iter_paragraphs(story):
                for run in paragraph.runs:
                    set_fonts(run)

    # Replace symbols that are vulnerable to WPS/PDF fallback. Mathematical display objects are retained.
    safe_replacements = {
        "BGA—MGA": "BGA与MGA",
        "BGA–MGA": "BGA与MGA",
        "Java—Python": "Java与Python",
        "Hodges–Lehmann": "Hodges-Lehmann",
        "E′": "E扩展",
        "A—D": "A至D",
        "—": "-",
        "–": "-",
        "→": "至",
        "−": "-",
        "×10⁻⁵": "e-5",
        "×10⁻⁸": "e-8",
        "×10⁻⁴": "e-4",
        "×10⁻⁷": "e-7",
        "×10⁻⁶": "e-6",
        "N₀": "N0",
        "∈": " in ",
        "≤": "不大于",
        "≥": "不小于",
        "√": "平方根",
        "ε": "epsilon",
        "Φ": "合法化函数",
        "κ": "排序键",
    }
    for paragraph in iter_paragraphs(document):
        replace_in_runs(paragraph, safe_replacements)

    # Make the Chinese keyword label an explicit, reliably rendered CJK run.
    set_single_run_text(
        document.paragraphs[5],
        "关键词：三相油浸式变压器；目录约束；约束优化；离散优化；遗传算法；双档案；工程精算一致性",
    )

    # Replace placeholder superscripts with ASCII while keeping the author block visibly pending.
    set_single_run_text(document.paragraphs[1], "作者姓名1，作者姓名2，作者姓名3*")

    # Correct ambiguous resource-consumption wording and define the input-protocol boundary.
    set_single_run_text(
        document.paragraphs[4],
        "针对三相油浸式变压器离散设计中材料目录离散、结构参数条件依赖及工程精算代价高的问题，提出一种目录约束双档案遗传优化方法。该方法以完整目录记录为原子决策单元，并通过合法化函数派生层数、油道和冷却关联变量，避免遗传操作产生无法由业务目录还原的虚拟规格。严格可行档案仅保存满足全部硬约束的候选，近可行档案保留可计算但存在约束超限的边界候选；覆盖初始化、耦合交叉和运行内已评价候选池注入构成默认增强搜索策略，诊断反馈局部变异作为可选扩展策略单独评估。基于五类预定义工程场景、每个场景30个配对随机种子的测试结果表明：相较目录约束基础遗传算法（BGA），在大容量长圆、散热器工况甲和工况乙中，MGA默认增强策略与BGA的严格可行发现率差异经双侧McNemar精确检验，并在5场景检验族内采用Holm校正后，p值分别为3.05e-5和1.86e-8；在双方共同严格成功的配对中，中容量长圆、波纹场景和大容量长圆、散热器工况甲的最低严格成本Hodges-Lehmann变化率分别为-0.83%和-7.02%。中小容量长圆、波纹场景的成本差异未达到统计显著。五个主场景中，BGA与MGA的第0代后真实精算调用中位数之比为1.85至4.82。结果支持：在当前冻结目录、价格快照与预算上限下，MGA默认增强策略可提高部分复杂配置的严格可行方案发现率；后续真实精算调用量较低仅为当前停止规则下的观测结果，不能替代首次严格解发现效率的判断。结论不外推为全部产品配置上的统一性能保证。",
    )
    set_single_run_text(
        document.paragraphs[7],
        "A catalogue-constrained dual-archive genetic optimization method is proposed for three-phase oil-immersed transformer design. Complete catalogue records are atomic decision units, and legalization derives layer, duct, and cooling variables to prevent virtual specifications that cannot be restored from business catalogues. Strict-feasible and near-feasible archives retain hard-constraint-satisfying and computable boundary candidates, respectively. Coverage initialization, coupled crossover, and run-time evaluated-pool injection form the default strategy; diagnostic local mutation is an optional extension. Five predefined engineering scenarios, each with 30 paired seeds, compare it with a catalogue-constrained baseline GA (BGA). In two large-capacity long-oval radiator scenarios, strict-feasible discovery-rate differences were tested by two-sided exact McNemar tests with five-scenario Holm adjustment (p=3.05e-5 and 1.86e-8). For jointly strict-feasible paired runs, Hodges-Lehmann cost-change estimates are -0.83% in the medium-capacity long-oval corrugated-tank scenario and -7.02% in large-capacity long-oval radiator scenario I; the small-to-medium-capacity long-oval corrugated-tank result is not significant. The ratio of median post-initialization engineering evaluations for BGA to MGA ranges from 1.85 to 4.82 across five scenarios. Under frozen catalogue and price snapshots and the same maximum actual-evaluation budget, the results support improved strict-feasible discovery for some of the tested complex configurations. Evaluation counts are descriptive under the stopping rule and do not measure first strict-feasible discovery efficiency; conclusions do not guarantee uniform performance.",
    )
    set_single_run_text(
        document.paragraphs[11],
        "为避免以构造数值替代工程证据，本文五个主比较场景的算法结果来自2026-09-24归档的真实运行轨迹；2026-09-28补充完成Java与Python一致性核验及泛化边界检查，二者仅用于界定计算链一致性与适用边界，不参与主比较统计。主比较覆盖五类来自实际业务系统的冻结工程配置快照（以下简称真实工程配置）：场景A为50 kVA圆形铁芯、波纹油箱，场景B为100 kVA长圆铁芯、波纹油箱，场景C为300 kVA长圆铁芯、波纹油箱，场景D为2 000 kVA长圆铁芯、散热器油箱工况甲，场景E为2 000 kVA长圆铁芯、散热器油箱工况乙。场景D与场景E的容量和主结构相同，但冻结的页面条件、目录快照与实际搜索规模不同，故作为独立场景分别报告。每个场景在BGA与MGA下采用相同的30个随机种子配对运行。另设一个200 kVA圆形铁芯、波纹油箱边界场景，双方均未得到严格解，仅用于边界检查，不进入主比较和显著性结论。",
    )
    set_single_run_text(
        document.paragraphs[12],
        "已有研究已将遗传算法、启发式算法和多目标进化算法用于变压器结构与制造成本优化，国外代表性研究见文献[1]至文献[3]；国内研究也已将遗传、群智能和代理建模方法用于变压器优化设计，见文献[8]至文献[10]以及文献[16]。近年关于复杂约束进化优化的综述和方法梳理见文献[13]至文献[15]及文献[17]。但对完整目录记录、条件依赖和工程精算可复核性的显式处理仍较少。一类研究将几何尺寸作为连续或相互独立的离散变量，难以表达绝缘、材料和价格等绑定属性；另一类约束进化研究给出可行性规则、随机排序或epsilon约束，却不解决候选能否由工程目录唯一还原的问题，见文献[4]至文献[6]。本文从工程数据可还原性出发，构建目录约束的离散候选域，并将最终工程约束和近边界探索分别纳入双档案。",
    )
    set_single_run_text(
        document.paragraphs[18],
        "任一可搜索模块m的完整记录定义见式（1），其中t_m为导线或结构类别，id_m为目录记录标识，a_m为随记录绑定的几何、材料、绝缘、价格和工艺属性。候选x由铁芯、低压绕组、高压绕组、冷却记录及条件依赖变量构成，其中r_c、r_l、r_h和r_k为四个完整记录，z为由规则派生的变量向量。令phi(x,q,D)表示合法化过程；只有当完整记录可由目录D唯一还原、页面范围q满足且z=phi(x,q,D)时，候选才属于合法域Omega(q,D)。工程精算器E(x;q,p)在给定价格与计算参数快照p下返回性能、温升、重量和成本，完整优化问题由式（3）给出。",
    )
    set_single_run_text(
        document.paragraphs[60],
        "表4列出五个预定义主比较场景的真实工程配置快照。表中的第0代精算数N0是本次覆盖、历史和随机候选去重后实际调用精算器的数量；工作种群规模N按当前实现由N0确定。本文所称基准测试，是指在冻结场景、冻结预算和配对随机种子条件下的统一比较协议，并非公共数据集意义上的开放基准。本文不以未导出的组合计数替代真实范围，故报告实际N0而非推断性的“严格可行域规模”。第0代后真实精算总预算B'覆盖演化和末端细化的全部未命中缓存调用；200 kVA圆形、波纹边界场景的检查结果单列说明，不并入表4。",
    )
    set_single_run_text(
        document.paragraphs[64],
        "严格可行率的5个场景构成一个Holm校正族；满足共同成功数不少于6的3个成本比较构成另一个Holm校正族。对同一种子s的共同严格成功配对，成本变化率定义为r_s=(C_MGA,s-C_BGA,s)/C_BGA,s*100%。成本差异使用双侧Wilcoxon符号秩检验、Hodges-Lehmann估计和秩双列相关系数r_rb报告；Hodges-Lehmann估计的95%置信区间采用配对非参数bootstrap（20 000次重采样、固定随机种子）计算。各场景不按容量、候选数或共同成功样本量加权，跨场景叙述仅作描述性概括。所有参数在正式运行前冻结，未根据本轮结果回调。",
    )
    set_single_run_text(
        document.paragraphs[68],
        "图4给出五个预定义场景的严格可行发现率及第0代后真实精算调用中位数。严格可行发现的差异具有场景相关性：场景D和场景E中，MGA均达到100.0%严格成功，而BGA分别为40.0%与3.3%；场景C中双方均为100.0%；场景B和场景A中，MGA分别为80.0%与20.0%，相对BGA的60.0%与16.7%仅为有限差异。经5场景Holm校正后，只有场景D和场景E的严格可行率差异达到p<0.05。所有主场景的第0代均未产生严格可行候选，严格解均出现在后续搜索阶段。BGA与MGA的后续真实精算调用中位数之比为1.85至4.82；相应地，MGA的调用量约为BGA的20.7%至54.1%。该数值受停止规则与缓存复用共同影响，仅作为运行资源消耗的描述性结果，不能据此单独推断首次发现严格解更快。",
    )
    set_single_run_text(
        document.paragraphs[72],
        "表6和图5报告严格成功率与共同严格成功配对上的成本结果。场景C中双方30个种子均严格成功，MGA相对BGA的Hodges-Lehmann成本变化率为-0.83%（95% CI：-1.03%，-0.83%；Holm p=3.84e-5；r_rb=-0.885）。场景D的共同成功配对数为12，成本变化率为-7.02%（95% CI：-11.96%，-3.75%；Holm p=9.77e-4；r_rb=-1.000）。场景B的共同成功配对数为14，成本变化率为0.00%，未达到统计显著（Holm p=0.5176；r_rb=-0.286）。场景A与场景E的共同成功配对数分别为2和1，因样本不足不作成本推断。因而，本文仅将成本结果解释为共同严格成功条件下的配对经济性，不将严格成功率与成本中位数混合为统一优劣指标。",
    )

    # Attach equation numbers to their formula, rather than leave (6) and (7) as isolated paragraphs.
    # Cost is formula (4); archive ranking, diversity, and complexity follow as (5)-(8).
    replace_math_text(document.paragraphs[39], {"(4)": "(5)"})
    replace_math_text(document.paragraphs[41], {"(5)": "(6)"})
    append_math_number(document.paragraphs[47], 7)
    append_math_number(document.paragraphs[49], 8)
    set_single_run_text(
        document.paragraphs[40],
        "多样性采用四项指标：完整记录签名数、四模块唯一组合数、工作种群平均汉明距离和模块熵。完整记录签名由铁芯、低压、高压和冷却四个完整模块记录的联合编码构成；四模块唯一组合数统计当前种群中互不相同的四元组合。对于任意两个候选，分别比较铁芯、低压、高压和冷却四个模块：某模块的完整目录记录ID不同记为1，否则记为0，故汉明距离d_H的取值范围为[0,4]。平均汉明距离与模块熵由式（6）给出，其中p_m(v)为模块m取值v的频率。距离和熵只描述离散目录选择的多样性，不将目录外几何组合计入搜索空间。",
    )
    remove_paragraph(document.paragraphs[50])
    remove_paragraph(document.paragraphs[48])

    # Appendix/figure/table text after paragraph removal needs lookup by content.
    for paragraph in document.paragraphs:
        if paragraph.text.startswith("图6和表8给出独立扩展批次"):
            replace_in_runs(paragraph, {"E扩展": "E扩展", "A至D": "A至D"})
        if paragraph.text.startswith("表8 独立扩展批次"):
            replace_in_runs(paragraph, {"E扩展": "E扩展"})

    # Table 8 header is explicit about the comparison baseline.
    set_cell_text(document.tables[9].cell(0, 5), "去双档案\n（诊断扩展）")

    # Complete the older core references and add recent, verified related work.
    references = {
        115: "[1] J. W. Nims, R. E. Smith, and A. A. El-Keib, “Application of a genetic algorithm to power transformer design,” Electric Machines & Power Systems, vol. 24, no. 6, pp. 669-680, 1996. DOI: 10.1080/07313569608955702.",
        116: "[2] P. S. Georgilakis, M. A. Tsili, and A. T. Souflaris, “A heuristic solution to the transformer manufacturing cost optimization problem,” Journal of Materials Processing Technology, vol. 181, nos. 1-3, pp. 260-266, 2007. DOI: 10.1016/j.jmatprotec.2006.03.034.",
        117: "[3] E. I. Amoiralis, M. A. Tsili, D. G. Paparigas, and A. G. Kladas, “Global transformer design optimization using deterministic and nondeterministic algorithms,” IEEE Transactions on Industry Applications, vol. 50, no. 1, pp. 383-394, 2014. DOI: 10.1109/TIA.2013.2288417.",
        118: "[4] K. Deb, “An efficient constraint handling method for genetic algorithms,” Computer Methods in Applied Mechanics and Engineering, vol. 186, nos. 2-4, pp. 311-338, 2000. DOI: 10.1016/S0045-7825(99)00389-8.",
        119: "[5] T. P. Runarsson and X. Yao, “Stochastic ranking for constrained evolutionary optimization,” IEEE Transactions on Evolutionary Computation, vol. 4, no. 3, pp. 284-294, 2000. DOI: 10.1109/4235.873238.",
        120: "[6] C. A. C. Coello, “Theoretical and numerical constraint-handling techniques used with evolutionary algorithms: a survey of the state of the art,” Computer Methods in Applied Mechanics and Engineering, vol. 191, nos. 11-12, pp. 1245-1287, 2002. DOI: 10.1016/S0045-7825(01)00323-1.",
        121: "[7] A. E. Eiben and J. E. Smith, Introduction to Evolutionary Computing, 2nd ed. Berlin, Germany: Springer, 2015. DOI: 10.1007/978-3-662-44874-8.",
    }
    for index, text in references.items():
        set_single_run_text(document.paragraphs[index - 2], text)  # indices shift after deleting two formula-label paragraphs

    added = [
        "[13] J. Liang, X. Ban, K. Yu, et al., “A survey on evolutionary constrained multiobjective optimization,” IEEE Transactions on Evolutionary Computation, vol. 27, no. 2, pp. 201-221, 2023. DOI: 10.1109/TEVC.2022.3155533.",
        "[14] C. A. C. Coello Coello, “Constraint-handling techniques used with evolutionary algorithms,” in Proceedings of the Genetic and Evolutionary Computation Conference Companion, Boston, MA, USA, 2022, pp. 1310-1333. DOI: 10.1145/3520304.3533640.",
        "[15] J. Liang, H. Lin, C. Yue, et al., “Evolutionary constrained multi-objective optimization: a review,” Vicinagearth, vol. 1, art. no. 5, 2024. DOI: 10.1007/s44336-024-00006-5.",
        "[16] Z. Li, L. Zhang, X. Chen, W. Xiao, M. Li, and B. Shi, “Three-phase transformer optimization design based on NSGA2 algorithm,” Journal of Physics: Conference Series, vol. 2591, no. 1, art. no. 012024, 2023. DOI: 10.1088/1742-6596/2591/1/012024.",
        "[17] 陈少淼, 陈瑞, 梁伟, 李仁发, 李智勇. 面向复杂约束优化问题的进化算法综述[J]. 软件学报, 2023, 34(2):565-581. DOI:10.13328/j.cnki.jos.006711.",
    ]
    for text in added:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.first_line_indent = Pt(-16)
        paragraph.paragraph_format.left_indent = Pt(16)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 0.88
        run = paragraph.add_run(text)
        set_fonts(run)
        run.font.size = Pt(6.5)

    # Keep all reference entries at the existing compact, readable style.
    ref_start = next(i for i, p in enumerate(document.paragraphs) if p.text == "参考文献") + 1
    for paragraph in document.paragraphs[ref_start:]:
        paragraph.paragraph_format.first_line_indent = Pt(-16)
        paragraph.paragraph_format.left_indent = Pt(16)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 0.88
        for run in paragraph.runs:
            set_fonts(run)
            run.font.size = Pt(6.5)

    # Validate that old superscript p-values / extension primes are gone from ordinary document text.
    remaining_text = "\n".join(p.text for p in document.paragraphs)
    forbidden = ["E′", "⁻", "−", "∪", "∈", "≤", "≥"]
    unexpected = [symbol for symbol in forbidden if symbol in remaining_text]
    if unexpected:
        raise RuntimeError(f"glyph-risk characters remain in ordinary text: {unexpected}")

    document.save(TARGET)
    print(TARGET)


if __name__ == "__main__":
    main()

"""V6.4：为双栏单栏排版重绘图 3—图 8 的高可读版本。

仅重绘图内文字与布局；数据、图题、正文结论和既有图片尺寸均不改动。
图内不重复放置长标题，避免与正文图题竞争版面。
"""
from pathlib import Path

import json
import matplotlib.pyplot as plt
import numpy as np
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(r"D:\IdeaProject\faladi\three_phase_ga_optimizer")
SOURCE = ROOT / "论文" / "目录约束双档案遗传优化三相变压器_版式与引用优化版V6.3-送外审样式与数学符号统一版.docx"
OUTPUT = ROOT / "论文" / "目录约束双档案遗传优化三相变压器_版式与引用优化版V6.4-送外审版式与图表可读性终校版.docx"
ASSET_DIR = ROOT / "论文" / "图表-送外审终校V6.4"
DATA_ROOT = ROOT / "测试归档" / "2026-09-24测试" / "论文真实数据"

INK = "#19324D"
BLUE = "#2E5F89"
TEAL = "#167C74"
ORANGE = "#BE7A2F"
PALE_BLUE = "#F4F8FC"
PALE_TEAL = "#EEF8F6"
PALE_ORANGE = "#FFF7EE"
GRID = "#D9E3EC"

plt.rcParams.update({
    "font.sans-serif": ["Microsoft YaHei", "SimHei", "Arial Unicode MS"],
    "axes.unicode_minus": False,
    "font.size": 16,
})


def save(fig, path):
    fig.savefig(path, dpi=420, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def box(ax, xy, width, height, title, subtitle, face, edge, title_size=16, subtitle_size=14):
    x, y = xy
    ax.add_patch(FancyBboxPatch(
        (x, y), width, height, boxstyle="round,pad=0.025,rounding_size=0.04",
        linewidth=1.7, edgecolor=edge, facecolor=face,
    ))
    ax.text(x + width / 2, y + height * .61, title, ha="center", va="center",
            fontsize=title_size, fontweight="bold", color=INK)
    ax.text(x + width / 2, y + height * .31, subtitle, ha="center", va="center",
            fontsize=subtitle_size, color="#42576B")


def arrow(ax, start, end, color=BLUE, rad=0):
    ax.add_patch(FancyArrowPatch(
        start, end, arrowstyle="-|>", mutation_scale=13, linewidth=1.9,
        color=color, connectionstyle=f"arc3,rad={rad}",
    ))


def figure3(path):
    """逻辑不交叉、适合单栏缩放的双档案循环图。"""
    fig, ax = plt.subplots(figsize=(6.4, 3.84))
    ax.set_xlim(0, 12); ax.set_ylim(0, 7); ax.axis("off")
    # 图题和注释由正文图题承担；图内只保留循环所必需的短标签，保证单栏缩小后仍可读。
    box(ax, (.30, 3.00), 2.45, 1.15, "候选工程精算", "性能、成本、约束", PALE_BLUE, BLUE, 13.6, 10.6)
    box(ax, (3.28, 4.62), 2.90, 1.03, "严格档案  A_f", "V(x)=0；按成本保留", PALE_TEAL, TEAL, 13.6, 10.6)
    box(ax, (3.28, 1.58), 2.90, 1.03, "近可行档案  A_n", "V(x)>0；按超限排序", PALE_ORANGE, ORANGE, 13.6, 10.6)
    box(ax, (7.20, 3.00), 2.15, 1.15, "父代抽样", "A_f 优先；A_n 可参与", "#FFFFFF", BLUE, 13.6, 9.9)
    box(ax, (10.02, 3.00), 1.65, 1.15, "新候选", "交叉、变异", PALE_BLUE, BLUE, 13.2, 10.1)
    box(ax, (6.70, .25), 3.55, .83, "评价与更新", "档案、候选池与种群更新", "#F8FAFC", TEAL, 13.6, 9.9)
    arrow(ax, (2.75, 3.77), (3.28, 5.13))
    arrow(ax, (2.75, 3.38), (3.28, 2.09))
    arrow(ax, (6.18, 5.13), (7.20, 3.77))
    arrow(ax, (6.18, 2.09), (7.20, 3.38), color=ORANGE, rad=-.14)
    arrow(ax, (9.35, 3.58), (10.02, 3.58))
    arrow(ax, (10.85, 3.00), (10.22, 1.08), color=TEAL, rad=-.12)
    arrow(ax, (6.70, .66), (2.75, .66), color=TEAL)
    arrow(ax, (2.75, .66), (2.75, 3.00), color=TEAL)
    fig.tight_layout(pad=.2)
    save(fig, path)


def figure4(path):
    scenes = ["A", "B", "C", "D", "E"]
    bga_success = [16.7, 60.0, 100.0, 40.0, 3.3]
    mga_success = [20.0, 80.0, 100.0, 100.0, 100.0]
    bga_calls = [2754.5, 7066.5, 10000.0, 8000.0, 8000.0]
    mga_calls = [663.0, 2220.5, 2074.0, 3580.5, 4314.0]
    fig, (a1, a2) = plt.subplots(1, 2, figsize=(6.4, 2.45), gridspec_kw={"wspace": .45})
    x = np.arange(5); width = .35
    for ax, left, right, ylabel, ymax in (
        (a1, bga_success, mga_success, "严格可行率（%）", 110),
        (a2, bga_calls, mga_calls, "后续真实精算调用中位数（次）", 11000),
    ):
        ax.bar(x-width/2, left, width, label="BGA", color="#AAB9C7")
        ax.bar(x+width/2, right, width, label="MGA", color=TEAL)
        ax.set_xticks(x, scenes, fontsize=16)
        ax.set_ylabel(ylabel, fontsize=16)
        ax.set_ylim(0, ymax)
        ax.tick_params(axis="y", labelsize=14)
        ax.grid(axis="y", color=GRID, linewidth=.6)
        ax.spines[["top", "right"]].set_visible(False)
    a1.legend(frameon=False, fontsize=15, ncol=2, loc="upper left")
    a2.legend(frameon=False, fontsize=15, ncol=2, loc="upper right")
    fig.tight_layout(pad=.35)
    save(fig, path)


def figure5(path):
    """HL 点估计与置信区间：删去图内重复长标题，放大单栏可读文字。"""
    rows = [
        ("场景 B", 14, 0.00, -2.83, 1.27, "0.00%", "Holm p=0.5176"),
        ("场景 C", 30, -0.83, -1.03, -0.83, "-0.83%", "Holm p=3.84e-5"),
        ("场景 D", 12, -7.02, -11.96, -3.75, "-7.02%", "Holm p=9.77e-4"),
    ]
    fig, ax = plt.subplots(figsize=(6.35, 3.75))
    y = [2, 1, 0]
    ax.axvline(0, color="#9BAEC1", linewidth=1.4, zorder=0)
    for yi, (name, n, est, lo, hi, label, p) in zip(y, rows):
        ax.hlines(yi, lo, hi, color=TEAL, linewidth=2.8, zorder=2)
        ax.vlines([lo, hi], yi-.07, yi+.07, color=TEAL, linewidth=2.1, zorder=2)
        ax.scatter([est], [yi], s=58, color="#0D5C78", zorder=3)
        ax.text(-12.65, yi+.07, name, ha="left", va="center", fontsize=15, fontweight="bold", color=INK)
        ax.text(-12.65, yi-.20, f"共同严格成功 n={n}", ha="left", va="center", fontsize=12.5, color="#607388")
        ax.text(est+.28, yi+.13, label, ha="left", va="center", fontsize=14, fontweight="bold", color=INK)
        ax.text(est+.28, yi-.16, p, ha="left", va="center", fontsize=11.8, color="#536578")
    ax.set_xlim(-13.0, 2.0); ax.set_ylim(-.5, 2.75); ax.set_yticks([])
    ax.set_xticks([-12, -9, -6, -3, 0]); ax.tick_params(axis="x", labelsize=13, colors="#455A70")
    ax.set_xlabel("Hodges–Lehmann 成本变化率（%）", fontsize=15, labelpad=8)
    ax.text(-12.9, 2.54, "负值表示 MGA 成本较低；点为 HL 估计，横线为 95% CI。", fontsize=11.8, color="#586B7E")
    ax.spines[["top", "right", "left"]].set_visible(False)
    ax.spines["bottom"].set_color("#62768B")
    ax.grid(axis="x", color=GRID, linewidth=.65, zorder=0)
    fig.tight_layout(pad=.35)
    save(fig, path)


def figure6(path):
    labels = ["基础\nBGA", "双档案\n+耦合交叉", "+候选池\n注入", "+诊断\n探测", "去双档案\n（诊断扩展）"]
    data = {"A": [16.7, 20.0, 20.0, 43.3, 63.3], "B": [60.0, 63.3, 80.0, 96.7, 70.0], "D": [40.0, 100.0, 100.0, 100.0, 100.0]}
    colors = {"A": ORANGE, "B": BLUE, "D": TEAL}
    fig, ax = plt.subplots(figsize=(6.4, 2.55))
    for name, values in data.items():
        ax.plot(range(5), values, marker="o", markersize=6.4, linewidth=2.2,
                color=colors[name], label=f"场景 {name}")
    ax.set_xticks(range(5), labels, fontsize=13)
    ax.set_ylim(0, 112); ax.set_yticks([0, 50, 100])
    ax.set_ylabel("严格可行率（%）", fontsize=15); ax.tick_params(axis="y", labelsize=13)
    ax.grid(axis="y", color=GRID, linewidth=.6); ax.spines[["top", "right"]].set_visible(False)
    ax.legend(frameon=False, fontsize=13, ncol=3, loc="upper left")
    fig.tight_layout(pad=.35)
    save(fig, path)


def figure7(path):
    first = json.loads((DATA_ROOT / "03_主比较_BGA_MGA" / "diversity_failure_summary.json").read_text(encoding="utf-8"))
    second = json.loads((DATA_ROOT / "03_主比较_BGA_MGA" / "diversity_failure_runs_t2rad_runs_t2r429.json").read_text(encoding="utf-8"))
    rows = first["diversity"] + second["diversity"]
    code = ["S-429", "S-387", "S-409", "S-469", "S-487"]
    names = ["A", "B", "C", "D", "E"]
    fig, axes = plt.subplots(1, 2, figsize=(6.4, 2.5), gridspec_kw={"wspace": .58})
    configs = [("final_distance", "平均汉明距离（0—4）", "工作种群平均汉明距离"),
               ("signature_median", "严格档案签名数（个）", "严格档案结构覆盖")]
    for ax, (field, ylabel, label) in zip(axes, configs):
        relevant = [r for r in rows if field in r]
        bga = [next((r[field] for r in relevant if r["scene"] == s and r["method"] == "BGA"), np.nan) for s in code]
        mga = [next((r[field] for r in relevant if r["scene"] == s and r["method"] == "MGA"), np.nan) for s in code]
        x = np.arange(5); w = .36
        ax.bar(x-w/2, bga, w, color="#AAB9C7", label="BGA")
        ax.bar(x+w/2, mga, w, color=TEAL, label="MGA")
        ax.set_xticks(x, [f"{n}" for n in names], fontsize=13)
        ax.set_ylabel(ylabel, fontsize=14); ax.tick_params(axis="y", labelsize=12)
        ax.set_title(label, fontsize=14, loc="left", fontweight="bold")
        ax.grid(axis="y", color=GRID, linewidth=.6); ax.spines[["top", "right"]].set_visible(False)
    axes[0].set_ylim(0, 4.7); axes[1].set_ylim(0, 15)
    axes[0].legend(frameon=False, fontsize=12, ncol=2, loc="upper left")
    fig.tight_layout(pad=.35)
    save(fig, path)


def figure8(path):
    rows = [
        ("A BGA\n(n=25)", {"负载损耗":23,"空载损耗":1,"其他":1}),
        ("A MGA\n(n=24)", {"负载损耗":24}),
        ("B BGA\n(n=12)", {"空载损耗":6,"阻抗":3,"负载损耗":2,"其他":1}),
        ("B MGA\n(n=6)", {"负载损耗":5,"空载损耗":1}),
        ("D BGA\n(n=18)", {"空载损耗":7,"散热器中心距":5,"温升/电流":5,"阻抗":1}),
        ("E BGA\n(n=29)", {"散热器中心距":10,"温升/电流":11,"负载损耗":5,"空载损耗":1,"其他":2}),
    ]
    cats = ["负载损耗", "空载损耗", "阻抗", "散热器中心距", "温升/电流", "其他"]
    colors = {"负载损耗":BLUE,"空载损耗":TEAL,"阻抗":ORANGE,"散热器中心距":"#7C3AED","温升/电流":"#C94343","其他":"#6B7280"}
    fig, ax = plt.subplots(figsize=(6.4, 2.7)); y = np.arange(len(rows)); left = np.zeros(len(rows))
    for cat in cats:
        values = np.array([r[1].get(cat, 0) / sum(r[1].values()) * 100 for r in rows])
        ax.barh(y, values, left=left, height=.62, color=colors[cat], label=cat)
        left += values
    ax.set_yticks(y, [r[0] for r in rows], fontsize=12.5); ax.invert_yaxis()
    ax.set_xlim(0, 100); ax.set_xlabel("失败运行中主导约束构成（%）", fontsize=14)
    ax.tick_params(axis="x", labelsize=12); ax.grid(axis="x", color=GRID, linewidth=.6)
    ax.spines[["top", "right"]].set_visible(False)
    ax.legend(frameon=False, fontsize=11.8, ncol=3, loc="upper center", bbox_to_anchor=(.5, -0.28))
    fig.tight_layout(pad=.35)
    save(fig, path)


def replace_image_blob(doc, paragraph_index, image_path):
    paragraph = doc.paragraphs[paragraph_index]
    rel_ids = paragraph._p.xpath(".//a:blip/@r:embed")
    if len(rel_ids) != 1:
        raise RuntimeError(f"paragraph {paragraph_index}: expected one image, got {rel_ids}")
    doc.part.related_parts[rel_ids[0]]._blob = image_path.read_bytes()


def set_math_font(run, size=7.5):
    run.font.name = "Cambria Math"
    run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    rfonts.set(qn("w:ascii"), "Cambria Math")
    rfonts.set(qn("w:hAnsi"), "Cambria Math")
    rfonts.set(qn("w:eastAsia"), "Cambria Math")


def restore_scientific_notation(doc):
    """将表 7 的程序式 e-5/e-8 恢复为排版用科学计数法。"""
    table = doc.tables[8]
    for row_index, source_token, mantissa, exponent in ((4, "3.05e-5", "3.05×10", "-5"), (5, "1.86e-8", "1.86×10", "-8")):
        p = table.cell(row_index, 1).paragraphs[0]
        old = p.text
        if source_token not in old:
            continue
        # 保留 b/c 与分号等原有文字，只替换科学计数法本体。
        before, after = old.split(source_token, 1)
        for child in list(p._p):
            if child.tag != qn("w:pPr"):
                p._p.remove(child)
        p.paragraph_format.first_line_indent = Pt(0)
        p.alignment = 1
        r = p.add_run(before); r.font.name = "Times New Roman"; r.font.size = Pt(7.5)
        m = p.add_run(mantissa); set_math_font(m)
        e = p.add_run(exponent); set_math_font(e); e.font.superscript = True
        if after:
            r = p.add_run(after); r.font.name = "Times New Roman"; r.font.size = Pt(7.5)


def main():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    images = {
        42: (ASSET_DIR / "图3_高可读循环图.png", figure3),
        67: (ASSET_DIR / "图4_严格可行率与精算调用量.png", figure4),
        71: (ASSET_DIR / "图5_HL成本变化率.png", figure5),
        77: (ASSET_DIR / "图6_预定义策略变体.png", figure6),
        79: (ASSET_DIR / "图7_多样性统计.png", figure7),
        84: (ASSET_DIR / "图8_失败运行主导约束.png", figure8),
    }
    for output, renderer in images.values():
        renderer(output)
    doc = Document(SOURCE)
    for index, (output, _) in images.items():
        replace_image_blob(doc, index, output)
    restore_scientific_notation(doc)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()

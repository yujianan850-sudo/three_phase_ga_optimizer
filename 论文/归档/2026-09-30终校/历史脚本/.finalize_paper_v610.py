"""V6.10：保留图4的稳定单栏版式，补全表1核心数学对象。"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = Path(r"D:\IdeaProject\faladi\three_phase_ga_optimizer")
SOURCE = ROOT / "论文" / "目录约束双档案遗传优化三相变压器_版式与引用优化版V6.7-送外审数学对象与图4终校版.docx"
OUTPUT = ROOT / "论文" / "目录约束双档案遗传优化三相变压器_版式与引用优化版V6.10-送外审数学对象与图4终校版.docx"


def set_font(run, size=7.5):
    run.font.size = Pt(size)
    run.font.name = "Times New Roman"
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr, value in (("ascii", "Times New Roman"), ("hAnsi", "Times New Roman"),
                        ("eastAsia", "宋体"), ("cs", "Times New Roman")):
        fonts.set(qn(f"w:{attr}"), value)
    rpr.set(qn("w:hint"), "eastAsia")


def clear_paragraph(p):
    for child in list(p._p):
        if child.tag != qn("w:pPr"):
            p._p.remove(child)


def m_run(value):
    run = OxmlElement("m:r")
    rpr = OxmlElement("m:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), "Cambria Math")
    rfonts.set(qn("w:hAnsi"), "Cambria Math")
    rpr.append(rfonts)
    run.append(rpr)
    node = OxmlElement("m:t")
    node.text = value
    run.append(node)
    return run


def math_plain(p, value):
    obj = OxmlElement("m:oMath")
    obj.append(m_run(value))
    p._p.append(obj)


def math_sub(p, base, sub):
    obj = OxmlElement("m:oMath")
    ssub = OxmlElement("m:sSub")
    base_node = OxmlElement("m:e")
    base_node.append(m_run(base))
    sub_node = OxmlElement("m:sub")
    sub_node.append(m_run(sub))
    ssub.append(base_node)
    ssub.append(sub_node)
    obj.append(ssub)
    p._p.append(obj)


def set_cell_math(cell, parts):
    p = cell.paragraphs[0]
    clear_paragraph(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for kind, values in parts:
        if kind == "text":
            run = p.add_run(values)
            set_font(run)
        elif kind == "plain":
            math_plain(p, values)
        elif kind == "sub":
            math_sub(p, *values)


def main():
    doc = Document(SOURCE)
    table = doc.tables[0]
    set_cell_math(table.cell(3, 2), [("sub", ("g", "i")), ("text", " / "), ("sub", ("h", "j"))])
    set_cell_math(table.cell(4, 0), [("sub", ("d̄", "H"))])
    set_cell_math(table.cell(4, 2), [("sub", ("H", "m"))])
    set_cell_math(table.cell(4, 4), [("plain", "B′")])
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()

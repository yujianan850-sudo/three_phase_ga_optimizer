"""Explicit illustrative author metadata and minimal terminology polishing."""
from pathlib import Path
from copy import deepcopy
import json, hashlib, sys
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

BASE=Path(__file__).resolve().parent
ROOT=BASE.parent.parent
SOURCE=next((ROOT/'归档/2026-10-01投稿准备/历史稿件').glob('*_送审修复版.docx'))
OUTPUT=ROOT/SOURCE.name.replace('_送审修复版','_作者占位优化版')
if OUTPUT.exists(): raise FileExistsError(OUTPUT)
doc=Document(SOURCE)

def replace_text(p, old, new):
    nodes=p._p.xpath('.//w:t')
    combined=''.join(n.text or '' for n in nodes)
    start=combined.find(old)
    if start<0: return False
    end=start+len(old); offset=0; inserted=False
    for node in nodes:
        value=node.text or ''; left=offset; right=offset+len(value); offset=right
        if right<=start or left>=end: continue
        prefix=value[:max(0,start-left)]
        suffix=value[max(0,end-left):] if end<right else ''
        node.text=prefix+(new if not inserted else '')+suffix
        inserted=True
    return True

def reset(p, content, size):
    p.clear()
    r=p.add_run(content);r.font.name='Times New Roman';r.font.size=Pt(size)
    r._r.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')

author=next(p for p in doc.paragraphs if p.text.startswith('作者姓名1'))
reset(author,'甲、乙、丙、丁*（作者占位）',9)
meta=next(p for p in doc.paragraphs if 'xxx@xxx.com' in p.text)
reset(meta,'（示例单位，示例城市；*通信作者：丁；E-mail：ding@example.com；基金项目：待补充）',8.5)

changes=[
 ('最低严格成本标准化Hodges–Lehmann位移率','最低严格成本的标准化Hodges–Lehmann位移'),
 ('当前主比较测试支持BGA与MGA默认增强策略的主比较','本研究已完成BGA与MGA默认增强策略的主比较'),
 ('独立扩展配置E扩展中全部变体均未获得严格解','场景E扩展中全部变体均未获得严格解'),
 ('五个真实主场景的冻结搜索快照与可复核信息','五个主比较场景的冻结搜索快照与可复核信息'),
]
for old,new in changes:
    count=sum(replace_text(p,old,new) for p in doc.paragraphs)
    assert count==1,(old,count)

doc.save(OUTPUT)
original=Document(SOURCE);final=Document(OUTPUT)
assert len(final.tables)==14
assert [t._tbl.xml for t in final.tables]==[t._tbl.xml for t in original.tables]
assert [r.blob for r in final.part.related_parts.values() if hasattr(r,'image')]==[r.blob for r in original.part.related_parts.values() if hasattr(r,'image')]
assert [p.text for p in final.paragraphs if p.text.startswith('[')]==[p.text for p in original.paragraphs if p.text.startswith('[')]
assert final.element.xpath('.//m:oMath') and len(final.element.xpath('.//m:oMath'))==len(original.element.xpath('.//m:oMath'))
report={'output':str(OUTPUT),'author_metadata':'explicit illustrative placeholders; replace before submission','changes':changes,'all_14_tables_unchanged':True,'all_images_unchanged':True,'references_unchanged':True,'sha256':hashlib.sha256(OUTPUT.read_bytes()).hexdigest()}
(BASE/'作者占位优化核验.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False,indent=2))

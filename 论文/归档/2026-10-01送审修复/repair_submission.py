"""Minimal manuscript repairs; preserve prior manuscript and measured results."""
from pathlib import Path
from copy import deepcopy
import sys,re,json,hashlib
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

BASE=Path(__file__).resolve().parent
ROOT=next(p for p in BASE.parents if (p/'ga').is_dir())
SOURCE=next((ROOT/'论文/归档/2026-10-01投稿准备/历史稿件').glob('*_指标口径修订版.docx'))
OUTPUT=ROOT/'论文'/SOURCE.name.replace('_指标口径修订版','_送审修复版')
if OUTPUT.exists() and '--retry-generated' not in sys.argv:raise FileExistsError(OUTPUT)
doc=Document(SOURCE)

def text(p,s,size=9):
    run=p.add_run(s);run.font.name='Times New Roman';run.font.size=Pt(size)
    run._r.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
    return run
def mr(s,size=9):
    r=OxmlElement('m:r')
    if re.search(r'[\u4e00-\u9fff]',s):
        mp=OxmlElement('m:rPr');normal=OxmlElement('m:nor');normal.set(qn('m:val'),'1');mp.append(normal);r.append(mp)
    pr=OxmlElement('w:rPr');fonts=OxmlElement('w:rFonts')
    for k in ('ascii','hAnsi','cs'):fonts.set(qn('w:'+k),'Cambria Math')
    fonts.set(qn('w:eastAsia'),'宋体')
    pr.append(fonts)
    for k in ('sz','szCs'):
        sz=OxmlElement('w:'+k);sz.set(qn('w:val'),str(round(size*2)));pr.append(sz)
    r.append(pr);t=OxmlElement('m:t');t.text=s;r.append(t);return r
def sub(a,b):
    s=OxmlElement('m:sSub');e=OxmlElement('m:e');e.append(mr(a));i=OxmlElement('m:sub');i.append(mr(b));s.extend([e,i]);return s
def sup(a,b):
    s=OxmlElement('m:sSup');e=OxmlElement('m:e');e.append(mr(a));i=OxmlElement('m:sup');i.append(mr(b));s.extend([e,i]);return s
def om(p,items):
    obj=OxmlElement('m:oMath')
    for item in items:obj.append(mr(item) if isinstance(item,str) else item)
    p._p.append(obj)
def find(prefix):return next(p for p in doc.paragraphs if p.text.startswith(prefix))
def visible(el):return ''.join(el.xpath('.//w:t/text()|.//m:t/text()'))
def display(p,rows,num):
    p.clear();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent=Pt(0)
    array=OxmlElement('m:eqArr');array.append(OxmlElement('m:eqArrPr'))
    for row in rows:
        e=OxmlElement('m:e')
        for node in row:e.append(mr(node) if isinstance(node,str) else node)
        array.append(e)
    om(p,[array]);text(p,'  ('+str(num)+')')

# Restore the accidentally dropped sizing expression as native math.
p=find('表4列出五个预定义')
before,after=p.text.split('工作种群规模N按当前实现由确定。',1)
p.clear();text(p,before+'工作种群规模按当前实现确定；本轮各配置取')
rad=OxmlElement('m:rad');pr=OxmlElement('m:radPr');hide=OxmlElement('m:degHide');hide.set(qn('m:val'),'1');pr.append(hide)
deg=OxmlElement('m:deg');e=OxmlElement('m:e');e.append(sub('N','0'));rad.extend([pr,deg,e])
om(p,['N = ⌈6',rad,'⌉'])
text(p,'，均未触发种群规模下限或第0代候选数上限。'+after)

# Replace only terminology, preserving all inline math objects and numerical values.
p=find('表6和图5报告')
for r in p.runs:r.text=r.text.replace('成本变化率','标准化HL成本位移')

# Equivalent wrapping of long equations, without shrinking body or table text.
p=next(p for p in doc.paragraphs if '(4)' in visible(p._p) and p._p.xpath('.//m:oMath'))
display(p,[['C(x;p) = ',sub('c','Fe'),sub('W','Fe'),' + ',sub('c','lv'),sub('W','lv'),' + ',sub('c','hv'),sub('W','hv')],
           ['+ ',sub('c','oil'),sub('W','oil'),' + ',sub('c','aux'),sub('W','aux')]],4)
p=next(p for p in doc.paragraphs if '(5)' in visible(p._p) and p._p.xpath('.//m:oMath'))
display(p,[['κ(x) = (V(x), ',sub('V','max'),'(x),'],[sub('n','v'),'(x), C(x))']],5)
p=next(p for p in doc.paragraphs if '(8)' in visible(p._p) and p._p.xpath('.//m:oMath'))
display(p,[['排序： O(N log N)'],['全量距离： O(',sup('N','2'),')']],8)

# Explicit run and structure-control fonts prevent LibreOffice default math sizing.
for obj in doc.element.xpath('.//m:oMath'):
    in_table=any(a.tag==qn('w:tc') for a in obj.iterancestors())
    size=15 if in_table else 18
    for r in obj.iter(qn('m:r')):
        pr=r.find(qn('w:rPr'))
        if pr is None:
            pr=OxmlElement('w:rPr');mathpr=r.find(qn('m:rPr'))
            r.insert(1 if mathpr is not None else 0,pr)
        fonts=pr.find(qn('w:rFonts'))
        if fonts is None:fonts=OxmlElement('w:rFonts');pr.insert(0,fonts)
        for k in ('ascii','hAnsi','cs'):fonts.set(qn('w:'+k),'Cambria Math')
        for k in ('sz','szCs'):
            el=pr.find(qn('w:'+k))
            if el is None:el=OxmlElement('w:'+k);pr.append(el)
            el.set(qn('w:val'),str(size))

# LibreOffice uses paragraph-mark formatting for the equation base size.
for p in doc.element.xpath('.//w:p[.//m:oMath]'):
    in_table=any(a.tag==qn('w:tc') for a in p.iterancestors())
    pp=p.find(qn('w:pPr'))
    if pp is None:pp=OxmlElement('w:pPr');p.insert(0,pp)
    rp=pp.find(qn('w:rPr'))
    if rp is None:rp=OxmlElement('w:rPr');pp.append(rp)
    for k in ('sz','szCs'):
        el=rp.find(qn('w:'+k))
        if el is None:el=OxmlElement('w:'+k);rp.append(el)
        el.set(qn('w:val'),'15' if in_table else '18')
    for pr in [el for el in obj.iter() if el.tag in {qn('m:'+k) for k in ('sSubPr','sSupPr','radPr','fPr','naryPr','eqArrPr','barPr')}]:
        ctrl=pr.find(qn('m:ctrlPr'))
        if ctrl is None:ctrl=OxmlElement('m:ctrlPr');pr.append(ctrl)
        rp=ctrl.find(qn('w:rPr'))
        if rp is None:rp=OxmlElement('w:rPr');ctrl.append(rp)
        for k in ('sz','szCs'):
            el=rp.find(qn('w:'+k))
            if el is None:el=OxmlElement('w:'+k);rp.append(el)
            el.set(qn('w:val'),str(size))

# Verify actual archives before writing each filename/hash pair.
archive=ROOT/'测试归档/2026-09-24测试/论文真实数据/02_场景筛选与预算预试验'
main=['runs_t2full.json','runs_t2rad.json','runs_t2r429.json']
variants=['runs_ablate.json','runs_ablate429.json','runs_ablate469.json','runs_ablate489.json']
hashes={name:hashlib.sha256((archive/name).read_bytes()).hexdigest().upper() for name in main+variants}
for row in doc.tables[12].rows:
    key=row.cells[0].text
    if key not in ('可复核性','策略变体归档标识'):continue
    p=row.cells[1].paragraphs[0];p.clear();p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    if key=='可复核性':
        text(p,'现存逐运行与统计JSON；主比较SQLite已清理，不能回放候选级轨迹。归档文件名及SHA-256前12位：',7.5)
        names=main
    else:
        text(p,'归档文件名及SHA-256前12位：',7.5);names=variants
    for name in names:
        run=text(p,'',7.5);run.add_break();text(p,name+'：'+hashes[name][:12],7.5)

doc.save(OUTPUT)
sys.path.insert(0,str(BASE.parent/'2026-10-01指标口径修订'))
from normalize_ooxml import normalize
import time
for attempt in range(5):
    try:normalize(OUTPUT);break
    except PermissionError:
        if attempt==4:raise
        time.sleep(0.3)
old=Document(SOURCE);new=Document(OUTPUT)
assert len(new.tables)==14 and len(new.inline_shapes)==9
for i in range(14):
    if i!=12:assert visible(old.tables[i]._tbl)==visible(new.tables[i]._tbl),(i,'table data changed')
assert [visible(p._p) for p in old.paragraphs if p.text.startswith('[')]==[visible(p._p) for p in new.paragraphs if p.text.startswith('[')]
assert '由确定' not in visible(new.element)
assert all(old.part.related_parts[a._inline.xpath('.//a:blip')[0].get(qn('r:embed'))]._blob == new.part.related_parts[b._inline.xpath('.//a:blip')[0].get(qn('r:embed'))]._blob for a,b in zip(old.inline_shapes,new.inline_shapes))
report={'source':str(SOURCE),'output':str(OUTPUT),'tables':14,'figures':9,'measured_tables_unchanged':True,'figures_unchanged':True,'archive_hashes':hashes,'author_metadata_pending':True}
(BASE/'修复核验.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps(report,ensure_ascii=False,indent=2))

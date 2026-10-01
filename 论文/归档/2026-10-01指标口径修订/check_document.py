"""Structural and statistical regression checks for the revised manuscript."""
from pathlib import Path
from collections import Counter
from copy import deepcopy
from hashlib import sha256
import json,re
from docx import Document
from docx.oxml.ns import qn

BASE=Path(__file__).resolve().parent
ROOT=next(p for p in BASE.parents if (p/'ga').is_dir())
src=next((ROOT/'论文').rglob('*_终校版.docx'))
dst=next((ROOT/'论文').rglob('*_指标口径修订版.docx'))
old=Document(src);doc=Document(dst)
def visible(node):return ''.join(node.xpath('.//w:t/text()|.//m:t/text()'))
for i in [1,2,3,4,5,6,7,8,9,10,11,13]:
    assert visible(old.tables[i]._tbl)==visible(doc.tables[i]._tbl),(i,'table data changed')
assert len(doc.inline_shapes)==9 and len(doc.tables)==14
refs=[p.text for p in doc.paragraphs if re.match(r'^\[\d+\]',p.text)]
assert len(refs)==17 and refs==[p.text for p in old.paragraphs if re.match(r'^\[\d+\]',p.text)]
body='\n'.join(visible(p._p) for p in doc.paragraphs if not re.match(r'^\[\d+\]',p.text))
citations=set()
for a,b in re.findall(r'文献\[(\d+)\]至文献\[(\d+)\]',body):citations.update(range(int(a),int(b)+1))
citations.update(map(int,re.findall(r'\[(\d+)\]',body)))
assert citations==set(range(1,18)),citations
figures=[visible(p._p) for p in doc.paragraphs if re.match(r'^图(?:\d|A1)',visible(p._p)) and p._p.xpath('.//w:fldSimple')]
tables=[visible(p._p) for p in doc.paragraphs if re.match(r'^表(?:\d|[ABC]1)',visible(p._p)) and p._p.xpath('.//w:fldSimple')]
assert [re.match(r'^图(\d+)',s).group(1) for s in figures if re.match(r'^图\d+',s)]==list(map(str,range(1,9)))
assert [re.match(r'^表(\d+)',s).group(1) for s in tables if re.match(r'^表\d+',s)]==list(map(str,range(1,10)))
assert len(figures)==9 and len(tables)==12
centered=[]
for i,p in enumerate(doc.paragraphs):
    if i in [0,1,6]:continue # title/author/title translation, not body prose
    if len(p.text)>90 and not p.text.startswith(('图','表')) and not p._p.xpath('.//m:oMath'):
        assert p.alignment!=1,(i,'long prose centered')
sizes=Counter()
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    if r.text.strip():
                        sizes[r.font.size.pt if r.font.size else None]+=1
                for size in p._p.xpath('.//m:oMath//w:sz'):
                    assert size.get(qn('w:val'))=='15'
assert set(sizes)=={7.5},sizes
alltext=visible(doc._element)
assert '[0,4]' not in alltext and '四模块平均汉明距离' not in alltext
assert not any(s in alltext for s in ['最终工作种群平均汉明距离','工作种群平均汉明距离和模块熵'])
assert '[0,5]' in alltext and '2026-09-28清理' in alltext
assert '作者姓名1' in alltext # deliberately retained; no fabricated identity
assert len(doc._element.xpath('.//w:fldSimple'))>=21
report=dict(tables=14,images=9,references=17,reference_round_trip=True,
    main_result_tables_unchanged=True,table_font_pt=7.5,figure_numbers='1-8, A1',
    table_numbers='1-9, A1, B1, C1',native_math_count=len(doc._element.xpath('.//m:oMath')),
    author_metadata_pending=True,docx_sha256=sha256(dst.read_bytes()).hexdigest())
(BASE/'文档回归核验.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps(report,ensure_ascii=False,indent=2))

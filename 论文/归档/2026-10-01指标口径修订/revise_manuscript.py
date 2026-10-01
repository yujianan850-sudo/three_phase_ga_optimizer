"""Targeted correction of diversity scope and retained-evidence wording.

Preserves source document, main statistics, all other images and tables.
"""
from pathlib import Path
from copy import deepcopy
import json,re
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

BASE=Path(__file__).resolve().parent
ROOT=next(p for p in BASE.parents if (p/'ga').is_dir())
SOURCE=next((ROOT/'论文/归档/2026-10-01投稿准备/历史稿件').glob('*_终校版.docx'))
OUTPUT=ROOT/'论文'/(SOURCE.stem.replace('_终校版','_指标口径修订版')+'.docx')
if OUTPUT.exists():
    raise FileExistsError(f'Output already exists: {OUTPUT}')
doc=Document(SOURCE)
paras=doc.paragraphs

def font(r,size=9):
    r.font.name='Times New Roman';r.font.size=Pt(size)
    r._r.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')

def mr(value,size=9):
    r=OxmlElement('m:r'); prop=OxmlElement('w:rPr')
    fonts=OxmlElement('w:rFonts')
    for k in ['ascii','hAnsi']:
        fonts.set(qn('w:'+k),'Cambria Math')
    prop.append(fonts)
    sz=OxmlElement('w:sz');sz.set(qn('w:val'),str(round(2*size)));prop.append(sz)
    r.append(prop);t=OxmlElement('m:t');t.text=value;r.append(t);return r

def sub(base,index,size=9):
    node=OxmlElement('m:sSub');e=OxmlElement('m:e');e.append(mr(base,size))
    ix=OxmlElement('m:sub');ix.append(mr(index,size));node.extend([e,ix]);return node

def bar(base,index,size=9):
    node=OxmlElement('m:bar');pr=OxmlElement('m:barPr');pos=OxmlElement('m:pos');pos.set(qn('m:val'),'top');pr.append(pos)
    e=OxmlElement('m:e');e.append(sub(base,index,size));node.extend([pr,e]);return node

def math(p,*nodes):
    obj=OxmlElement('m:oMath')
    for n in nodes:
        obj.append(mr(n) if isinstance(n,str) else n)
    p._p.append(obj)

def replace(p,value,body=True,size=9):
    p.clear();font(p.add_run(value),size)
    if body:
        p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent=Pt(18)
        p.paragraph_format.keep_with_next=False

def find(prefix):
    return next(p for p in doc.paragraphs if p.text.startswith(prefix))

replace(find('多样性采用四项指标'),
    '多样性采用完整记录签名数和五字段候选距离作描述性统计。完整记录签名由硅钢牌号、铁芯目录记录、低压导线记录、高压导线记录及冷却完整块的联合编码构成。候选距离对上述五字段逐项比较，不同记为1，相同记为0；铁芯牌号与铁芯记录分开计数，故d_H的范围为[0,5]，不是四模块距离。每代仅纳入已记录且完整可计算的候选，按固定步长抽样至多120条，计算式（6）的平均距离，K为样本数。图7(a)取同一场景、同一方法30次运行均有距离统计的最后一代，并报告其中位数；该代不是各运行的最终代，两方法端点也可能不同。图7(b)取各运行末个工作代严格档案的签名数中位数，不含末端细化后的档案更新。两项统计均不表示目录外几何组合的多样性。')

# Equation (6): use sample size K rather than the working-population symbol N.
p=paras[42];p.clear();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
frac=OxmlElement('m:f');num=OxmlElement('m:num');num.append(mr('2'));den=OxmlElement('m:den');den.append(mr('K(K−1)'));frac.extend([num,den])
summation=OxmlElement('m:nary');pr=OxmlElement('m:naryPr')
ch=OxmlElement('m:chr');ch.set(qn('m:val'),'∑');pr.append(ch)
lim=OxmlElement('m:limLoc');lim.set(qn('m:val'),'undOvr');pr.append(lim)
hide=OxmlElement('m:supHide');hide.set(qn('m:val'),'1');pr.append(hide)
summation.append(pr);ix=OxmlElement('m:sub');ix.append(mr('1≤i<j≤K'));sup=OxmlElement('m:sup');e=OxmlElement('m:e');e.append(sub('d','H'));e.append(mr('('));e.append(sub('x','i'));e.append(mr(','));e.append(sub('x','j'));e.append(mr(')'));summation.extend([ix,sup,e])
math(p,bar('d','H'),' = ',frac,summation)
font(p.add_run('   (6)'))

replace(find('图7呈现5个主场景'),
    '图7给出五个主场景的归档描述性多样性汇总。在各方法各自的末个共同覆盖代，BGA的五字段候选平均距离中位数均高于MGA；但两方法的统计代不同，候选集合也不是保留后的最终工作种群，因此不据此判断最终种群的多样性优劣。末个工作代严格档案的不重复签名数随场景和是否获得严格解而变化，未呈现单向优势。本文不把“多样性更高”作为MGA的实证结论。图8和表9按失败运行最小超限候选的主导约束归类，分母为未获得严格解的运行数，而非全部候选或评价次数；超限并列时按精算器固定输出顺序取首项，使每条失败运行仅归入一类。')
replace(find('图7 五个主场景'),
    '图7 五个主场景的五字段候选距离与末个工作代严格档案签名数（30次运行中位数；描述性汇总，统计端点见2.1）',False)

# Correct evidence availability rather than claiming missing files are reproducible.
p=find('本文贡献如下')
for r in p.runs:
    if '并保存可复核轨迹' in r.text:
        r.text=r.text.replace('并保存可复核轨迹','并留存逐运行及统计汇总记录')
p=find('主比较包含5个预定义场景')
for r in p.runs:
    r.text=r.text.replace('共300条完整运行轨迹','共300条逐运行汇总记录')
p=find('本文提出了面向三相油浸式变压器离散设计')
for r in p.runs:
    r.text=r.text.replace('每个场景30个配对随机种子的轨迹','每个场景30个配对随机种子的运行记录')
replace(find('表B1汇总五个实际业务系统'),
    '表B1汇总五个主比较场景的冻结输入边界，不表示全部页面字段或穷尽可行域。当前未建立独立snapshot UUID，以归档日期、场景及逐运行JSON标识运行。主比较原始SQLite已在2026-09-28清理，本次核查其300个原路径均为零字节占位；现有JSON支持重算成功计数与配对成本统计，不能恢复候选级轨迹或最终工作种群。图7只能核验历史汇总值及其提取口径。企业目录与价格不公开，外部复现仍需授权的脱敏配置和元数据。')
p=find('本文仍存在以下局限')
font(p.add_run('此外，主比较原始逐候选轨迹未留存，限制最终工作种群、多样性及首次严格解评价次数的再分析；上述分析需恢复备份或新增独立运行，不能由现有汇总补造。'))

# Table 1 retains the symbol/meaning pair positions and 7.5 pt style.
t=doc.tables[0]
replace(t.rows[4].cells[1].paragraphs[0],'五字段候选平均距离（0—5）',False,7.5)
p=t.rows[4].cells[2].paragraphs[0];p.clear();math(p,mr('K',7.5))
replace(t.rows[4].cells[3].paragraphs[0],'距离样本数，至多120',False,7.5)

t=doc.tables[12]
for row in t.rows:
    key=row.cells[0].text
    if key=='可复核性':
        replace(row.cells[1].paragraphs[0],
            '现存逐运行与统计JSON；主比较SQLite已清理，不能回放候选级轨迹。主运行JSON SHA-256前12位：A0F1266914A4、4D9E7D17C5C6、1030D418E12E。',False,7.5)

# Replace the Figure 7 blob only, preserving its existing size and paragraph.
shape=doc.inline_shapes[6]
blip=shape._inline.xpath('.//a:blip')[0]
rid=blip.get(qn('r:embed'))
doc.part.related_parts[rid]._blob=(BASE/'图7-五字段描述性统计.png').read_bytes()
shape._inline.docPr.set('descr','五字段候选距离：各方法末个共同覆盖代中位数；严格档案签名：末个工作代中位数。两方法距离端点可能不同，不是最终工作种群对比。')

# Equation (3) is an equivalent three-line native equation, avoiding narrow-column fragments.
p=paras[19];p.clear();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
array=OxmlElement('m:eqArr');apr=OxmlElement('m:eqArrPr');array.append(apr)
rows=[
    [sub('min','x∈Ω(q,D)'),' C(x;p)'],
    [sub('g','i'),'(E(x;q,p))≤0,  i∈',sub('I','g'),'(q)'],
    [sub('h','j'),'(E(x;q,p))=0,  j∈',sub('I','h'),'(q)']]
for nodes in rows:
    e=OxmlElement('m:e')
    for n in nodes:e.append(mr(n) if isinstance(n,str) else n)
    array.append(e)
math(p,array);font(p.add_run('  (3)'))

# Native subscript conversion for new prose and remaining ordinary symbolic runs.
mapping={'d_H':('d','H'),'t_m':('t','m'),'id_m':('id','m'),'a_m':('a','m'),
    'r_c':('r','c'),'r_l':('r','l'),'r_h':('r','h'),'r_k':('r','k'),
    'W_Fe':('W','Fe'),'W_lv':('W','lv'),'W_hv':('W','hv'),'W_oil':('W','oil'),'W_aux':('W','aux'),
    'c_Fe':('c','Fe'),'c_lv':('c','lv'),'c_hv':('c','hv'),'c_oil':('c','oil'),'c_aux':('c','aux'),
    'n_c':('n','c'),'n_l':('n','l'),'n_h':('n','h'),'n_k':('n','k')}
pattern=re.compile(r'(?<![A-Za-z0-9_])('+'|'.join(map(re.escape,sorted(mapping,key=len,reverse=True)))+r')(?![A-Za-z0-9_])')
for p in [paras[18],paras[22],paras[41],paras[50]]:
    for r in list(p.runs):
        if not pattern.search(r.text):continue
        cursor=0
        for match in pattern.finditer(r.text):
            if match.start()>cursor:
                cp=deepcopy(r._r)
                for el in list(cp):
                    if el.tag!=qn('w:rPr'):cp.remove(el)
                tx=OxmlElement('w:t');tx.text=r.text[cursor:match.start()];tx.set(qn('xml:space'),'preserve');cp.append(tx);r._r.addprevious(cp)
            obj=OxmlElement('m:oMath');obj.append(sub(*mapping[match.group()]));r._r.addprevious(obj);cursor=match.end()
        if cursor<len(r.text):
            cp=deepcopy(r._r)
            for el in list(cp):
                if el.tag!=qn('w:rPr'):cp.remove(el)
            tx=OxmlElement('w:t');tx.text=r.text[cursor:];tx.set(qn('xml:space'),'preserve');cp.append(tx);r._r.addprevious(cp)
        r._r.getparent().remove(r._r)

doc.save(OUTPUT)
from normalize_ooxml import normalize
normalize(OUTPUT)
check=Document(OUTPUT)
assert len(check.tables)==14 and len(check.inline_shapes)==9
assert len([p for p in check.paragraphs if re.match(r'^\[\d+\]',p.text)])==17
print(json.dumps(dict(output=str(OUTPUT),tables=14,images=9,references=17),ensure_ascii=False))

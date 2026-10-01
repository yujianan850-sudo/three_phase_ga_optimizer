"""Add live caption numbers and format the remaining display record tuple."""
from pathlib import Path
from copy import deepcopy
import re,sys
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from normalize_ooxml import normalize

path=Path(sys.argv[1]).resolve();doc=Document(path)
def mr(text):
    r=OxmlElement('m:r');pr=OxmlElement('w:rPr');f=OxmlElement('w:rFonts')
    for k in ['ascii','hAnsi']:f.set(qn('w:'+k),'Cambria Math')
    pr.append(f);sz=OxmlElement('w:sz');sz.set(qn('w:val'),'18');pr.append(sz);r.append(pr)
    t=OxmlElement('m:t');t.text=text;r.append(t);return r
def sub(base,index):
    s=OxmlElement('m:sSub');b=OxmlElement('m:e');b.append(mr(base));ix=OxmlElement('m:sub');ix.append(mr(index));s.extend([b,ix]);return s
for p in doc.paragraphs:
    match=re.match(r'^(图|表)([ABC]?)(\d+)(\s.*)$',p.text)
    if not match:continue
    category,appendix,number,tail=match.groups()
    pr=deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    p.clear()
    def text(s):
        r=p.add_run(s)
        if pr is not None:r._r.insert(0,deepcopy(pr))
    text(category+appendix)
    field=OxmlElement('w:fldSimple');key=('Figure' if category=='图' else 'Table')+('Appendix'+appendix if appendix else '')
    field.set(qn('w:instr'),' SEQ '+key+' \\* ARABIC ')
    r=OxmlElement('w:r')
    if pr is not None:r.append(deepcopy(pr))
    t=OxmlElement('w:t');t.text=number;r.append(t);field.append(r);p._p.append(field)
    text(tail)
p=doc.paragraphs[16];p.clear()
om=OxmlElement('m:oMath')
for node in [sub('r','m'),mr(' = ('),sub('t','m'),mr(', '),sub('id','m'),mr(', '),sub('a','m'),mr(')')]:om.append(node)
p._p.append(om);run=p.add_run('  (1)');run.font.name='Times New Roman';run.font.size=Pt(9)
# Table 1 denotes the mean distance, not the pairwise distance.
p=doc.tables[0].rows[4].cells[0].paragraphs[0];p.clear()
om=OxmlElement('m:oMath');bar=OxmlElement('m:bar');bp=OxmlElement('m:barPr')
pos=OxmlElement('m:pos');pos.set(qn('m:val'),'top');bp.append(pos)
e=OxmlElement('m:e');e.append(sub('d','H'));bar.extend([bp,e]);om.append(bar);p._p.append(om)
for size in p._p.xpath('.//w:sz'):size.set(qn('w:val'),'15')

# Older equation runs contained typed Vmax/nv rather than native subscripts.
pattern=re.compile(r'(?<![A-Za-z])(Vmax|nv)(?![A-Za-z])')
for node in list(doc.element.xpath('.//m:r')):
    texts=node.findall(qn('m:t'))
    if len(texts)!=1 or not pattern.search(texts[0].text or ''):continue
    content=texts[0].text;cursor=0
    for match in pattern.finditer(content):
        if match.start()>cursor:node.addprevious(mr(content[cursor:match.start()]))
        node.addprevious(sub(*({'Vmax':('V','max'),'nv':('n','v')}[match.group()])))
        cursor=match.end()
    if cursor<len(content):node.addprevious(mr(content[cursor:]))
    node.getparent().remove(node)

for p in doc.paragraphs:
    if p.text.startswith('设铁芯、低压绕组'):
        # Distinguish a full-set quadratic calculation from the actual capped sample.
        for run in p.runs:
            if '平均汉明距离的直接计算开销为O(N²)' in run.text:
                run.text=run.text.replace('平均汉明距离的直接计算开销为O(N²)，但其仅用于离线轨迹统计，不参与父代选择或候选生成。',
                    '全量平均汉明距离的直接计算为O(N²)；本轮离线统计以K≤120抽样，距离计算为O(K²)，不参与父代选择或候选生成。')
    if not p.text.startswith('表6和图5报告'):continue
    for run in list(p.runs):
        if 'rrb' not in run.text:continue
        segments=run.text.split('rrb')
        for j,seg in enumerate(segments):
            if seg:
                cp=deepcopy(run._r)
                for child in list(cp):
                    if child.tag!=qn('w:rPr'):cp.remove(child)
                t=OxmlElement('w:t');t.text=seg;t.set(qn('xml:space'),'preserve');cp.append(t);run._r.addprevious(cp)
            if j<len(segments)-1:
                obj=OxmlElement('m:oMath');obj.append(sub('r','rb'));run._r.addprevious(obj)
        run._r.getparent().remove(run._r)
doc.save(path);normalize(path)
print('Caption SEQ fields and native equation (1) added; visible numbering preserved.')
